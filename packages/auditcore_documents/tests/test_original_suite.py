"""Die 9 Originaltests aus audit_designer (tests/test_document_compare_service.py).

Unverändert in der Aussage, umgestellt auf die Kompatibilitätsfassade der
Bibliothek. Die Eingaben werden wie im Original mit python-docx erzeugt.
"""

from __future__ import annotations

from pathlib import Path
from time import perf_counter

import pytest

docx = pytest.importorskip("docx")

from auditcore_documents.legacy import DocumentCompareService  # noqa: E402
from auditcore_documents.settings import (  # noqa: E402
    DEFAULT_MEMO_OUTLINE,
    merge_settings,
    sanitise_settings,
)

Document = docx.Document


def _write_checklist(path: Path, questions: list[str]) -> None:
    document = Document()
    table = document.add_table(rows=0, cols=3)
    for question in questions:
        cells = table.add_row().cells
        cells[0].text = "☐"
        cells[1].text = question
        cells[2].text = ""
    document.save(path)


def _write_text(path: Path, absaetze: list[str]) -> None:
    document = Document()
    for absatz in absaetze:
        document.add_paragraph(absatz)
    document.save(path)


def test_checklist_compare_classifies_changed_removed_and_added(tmp_path: Path) -> None:
    old, new = tmp_path / "alt.docx", tmp_path / "neu.docx"
    _write_checklist(old, ["Frage unverändert", "Frage wird geändert", "Frage entfällt"])
    _write_checklist(new, ["Frage unverändert", "Frage wird deutlich geändert", "Neue Frage"])
    result = DocumentCompareService.compare(old, new, mode="checklist", threshold=70)
    assert result.mode == "checklist"
    assert (result.old_count, result.new_count) == (3, 3)
    assert (result.removed_count, result.added_count, result.changed_count) == (1, 1, 1)


def test_docx_render_contains_a_page_field_and_four_columns(tmp_path: Path) -> None:
    old, new = tmp_path / "alt.docx", tmp_path / "neu.docx"
    _write_checklist(old, ["Eine Frage"])
    _write_checklist(new, ["Eine neue Frage"])
    result = DocumentCompareService.compare(old, new, mode="checklist")
    result.metadata.update({"old_label": "Stand 2025", "new_label": "Stand 2026"})
    output = tmp_path / "result.docx"
    DocumentCompareService.render_docx(result, output, user="Test")
    document = Document(output)
    assert len(document.tables) == 1
    assert len(document.tables[0].columns) == 4
    headers = [cell.text for cell in document.tables[0].rows[0].cells]
    assert headers[1:3] == ["Stand 2025", "Stand 2026"]
    assert "NUMPAGES" in document.sections[0].footer._element.xml


def test_pdf_support_is_explicitly_text_mode(tmp_path: Path) -> None:
    pdf = tmp_path / "input.pdf"
    pdf.write_bytes(b"not a real pdf")
    try:
        DocumentCompareService.read(pdf)
    except Exception as exc:  # noqa: BLE001 - wie im Original
        assert "PDF" in str(exc) or "pypdf" in str(exc) or "pdftotext" in str(exc)


def test_personal_layout_overrides_department_and_is_sanitised() -> None:
    personal = sanitise_settings(
        {
            "output_profile": "memo",
            "memo_layout": {
                "header_text": "Mein\nBriefkopf",
                "font_family": "Arial",
                "body_font_size_pt": 99,
                "accent_color": "005EA8",
                "outline": ["Anlass", "Bewertung"],
            },
        }
    )
    effective, sources = merge_settings(
        {"memo_layout": {"font_family": "Hessen Gellix"}}, personal, {}
    )
    assert effective["memo_layout"]["header_text"] == "Mein\nBriefkopf"
    assert effective["memo_layout"]["font_family"] == "Arial"
    assert effective["memo_layout"]["body_font_size_pt"] == 18
    assert effective["memo_layout"]["accent_color"] == "#005EA8"
    assert effective["memo_layout"]["outline"][-1].endswith("{{vergleich}}")
    assert sources["memo_layout.font_family"] == "personal"


def test_memo_layout_is_used_by_docx_renderer(tmp_path: Path) -> None:
    old, new = tmp_path / "alt.docx", tmp_path / "neu.docx"
    _write_checklist(old, ["Alte Frage"])
    _write_checklist(new, ["Neue Frage"])
    result = DocumentCompareService.compare(old, new, mode="checklist", threshold=70)
    output = tmp_path / "styled.docx"
    DocumentCompareService.render_docx(
        result,
        output,
        user="Erika Muster",
        profile="memo",
        layout={
            "header_text": "Dienststelle\nVermerk",
            "font_family": "Arial",
            "body_font_size_pt": 11,
            "heading_font_size_pt": 16,
            "line_spacing": 1.2,
            "page_margin_cm": 2.5,
            "accent_color": "#005EA8",
            "footer_text": "Nur für den Dienstgebrauch",
            "show_page_numbers": False,
            "show_file_metadata": False,
            "outline": ["Anlass", "Feststellungen {{vergleich}}", "Bewertung"],
        },
    )
    document = Document(output)
    assert document.styles["Normal"].font.name == "Arial"
    assert "Dienststelle" in document.sections[0].header.paragraphs[0].text
    assert "Erika Muster" in document.sections[0].header.paragraphs[0].text
    assert "Nur für den Dienstgebrauch" in document.sections[0].footer.paragraphs[0].text
    assert "NUMPAGES" not in document.sections[0].footer._element.xml
    body_xml = document._element.body.xml
    assert body_xml.index("Feststellungen") < body_xml.index("<w:tbl")
    assert body_xml.index("<w:tbl") < body_xml.index("Bewertung")


def test_one_hundred_checklist_questions_finish_under_ten_seconds(tmp_path: Path) -> None:
    old, new = tmp_path / "alt.docx", tmp_path / "neu.docx"
    _write_checklist(old, [f"Prüffrage {index}" for index in range(100)])
    _write_checklist(
        new,
        [f"Prüffrage {i} ergänzt" if i % 10 == 0 else f"Prüffrage {i}" for i in range(100)],
    )
    started = perf_counter()
    result = DocumentCompareService.compare(old, new, mode="checklist", threshold=70)
    assert perf_counter() - started < 10
    assert (result.old_count, result.new_count) == (100, 100)
    assert DEFAULT_MEMO_OUTLINE


def test_tabellenzeilen_zerreissen_nicht_am_seitenumbruch(tmp_path: Path) -> None:
    from docx.oxml.ns import qn

    old, new = tmp_path / "alt.docx", tmp_path / "neu.docx"
    fragen = [
        "Wurde die Vergabe dokumentiert?",
        "Liegt der Zuwendungsbescheid vor?",
        "Wurde die Publizitaet geprueft?",
    ]
    _write_checklist(old, fragen)
    _write_checklist(new, ["Wurde die Vergabe vollstaendig dokumentiert?", *fragen[1:]])
    ergebnis = DocumentCompareService.compare(old, new, mode="checklist", threshold=70)
    ziel = tmp_path / "vergleich.docx"
    DocumentCompareService.render_docx(ergebnis, ziel, title="Probe")
    tabelle = Document(str(ziel)).tables[0]
    assert tabelle.rows
    for nummer, zeile in enumerate(tabelle.rows):
        eigenschaften = zeile._tr.find(qn("w:trPr"))
        assert eigenschaften is not None, f"Zeile {nummer} ohne trPr"
        assert eigenschaften.find(qn("w:cantSplit")) is not None


def test_verschobene_absaetze_sind_keine_streichung(tmp_path: Path) -> None:
    alt, neu = tmp_path / "alt.docx", tmp_path / "neu.docx"
    _write_text(
        alt,
        [
            "Die Pruefbehoerde prueft das Vorhaben.",
            "Der Beguenstigte legt die Belege vor.",
            "Die Auszahlung erfolgt danach.",
        ],
    )
    _write_text(
        neu,
        [
            "Der Beguenstigte legt die Belege vor.",
            "Die Pruefbehoerde prueft das Vorhaben.",
            "Die Auszahlung erfolgt danach.",
        ],
    )
    ergebnis = DocumentCompareService.compare(alt, neu, mode="text")
    assert (ergebnis.removed_count, ergebnis.added_count, ergebnis.moved_count) == (0, 0, 1)
    verschoben = [zeile for zeile in ergebnis.rows if zeile.status == "moved"]
    assert all("→" in zeile.location for zeile in verschoben)
    assert all(zeile.old_text == zeile.new_text for zeile in verschoben)


def test_echte_streichung_bleibt_streichung(tmp_path: Path) -> None:
    alt, neu = tmp_path / "alt.docx", tmp_path / "neu.docx"
    _write_text(
        alt,
        ["Die Pruefbehoerde prueft das Vorhaben.", "Eine Auftragsvergabe erfolgte beschraenkt."],
    )
    _write_text(neu, ["Die Pruefbehoerde prueft das Vorhaben."])
    ergebnis = DocumentCompareService.compare(alt, neu, mode="text")
    assert (ergebnis.removed_count, ergebnis.moved_count) == (1, 0)
