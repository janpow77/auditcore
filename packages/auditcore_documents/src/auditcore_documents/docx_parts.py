"""Bausteine der DOCX-Synopse: Zellformat, Wortmarkierung, Fassungszellen (python-docx).

Alle python-docx-Importe erfolgen verzögert (Extra ``docx-render``); die
Objekte des Originals werden unverändert erzeugt.
"""

from __future__ import annotations

import difflib
from dataclasses import dataclass
from datetime import datetime
from typing import Any

from auditcore_documents.model import CompareRow

NAVY = "14006E"
TINT = "E6EFFF"
HEADER_FILL = "BACBFF"
REMOVED = "B91C1C"
ADDED = "166534"


@dataclass(frozen=True)
class DocxStyle:
    """Aus dem Layout abgeleitete, begrenzte Gestaltungswerte (wie im Original)."""

    font_family: str
    body_size: float
    heading_size: float
    line_spacing: float
    margin_cm: float
    accent_color: str
    header_text: str
    footer_text: str
    show_page_numbers: bool
    show_file_metadata: bool

    @classmethod
    def from_layout(cls, layout: dict[str, Any], header: str | None) -> DocxStyle:
        """Ungültige Zahlen lösen wie im Original ``ValueError`` aus."""
        accent_color = str(layout.get("accent_color") or NAVY).lstrip("#").upper()
        if len(accent_color) != 6 or any(value not in "0123456789ABCDEF" for value in accent_color):
            accent_color = NAVY
        return cls(
            font_family=str(layout.get("font_family") or "Hessen Gellix")[:80],
            body_size=max(8.0, min(18.0, float(layout.get("body_font_size_pt", 10)))),
            heading_size=max(10.0, min(24.0, float(layout.get("heading_font_size_pt", 14)))),
            line_spacing=max(1.0, min(2.0, float(layout.get("line_spacing", 1)))),
            margin_cm=max(1.0, min(4.0, float(layout.get("page_margin_cm", 2)))),
            accent_color=accent_color,
            header_text=str(
                header if header is not None else layout.get("header_text") or "Vermerk"
            ),
            footer_text=str(layout.get("footer_text") or ""),
            show_page_numbers=bool(layout.get("show_page_numbers", True)),
            show_file_metadata=bool(layout.get("show_file_metadata", True)),
        )


def format_timestamp(value: object) -> str:
    """Zeitangabe in der lokalen Zeitzone des Prozesses (wie im Original)."""
    if not value:
        return "unbekannt"
    try:
        if isinstance(value, (int, float)):
            parsed = datetime.fromtimestamp(value)
        else:
            parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return parsed.astimezone().strftime("%d.%m.%Y, %H:%M")
    except (ValueError, TypeError, OSError):
        return str(value)


def keep_row_together(row: Any) -> None:
    """Verhindert, dass eine Tabellenzeile am Seitenumbruch zerrissen wird.

    Ohne dieses Flag brach Word die Zeile mitten im Satz um: Die Fundstelle
    stand dann auf der einen Seite, die Hälfte des Textes auf der nächsten,
    und die Fundstellenspalte der Folgeseite blieb leer. In einem Vermerk
    ist eine Textstelle ohne zugehörige Fundstelle wertlos.

    Ist eine Zeile höher als eine Seite, teilt Word sie trotzdem — dieser
    Rückfall ist gewollt, sonst ginge der Inhalt ganz verloren.
    """
    add_flag(row._tr.get_or_add_trPr(), "w:cantSplit")


def add_flag(properties: Any, tag: str) -> None:
    """Hängt ``<tag w:val="true"/>`` an ein Eigenschaftselement an."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    flag = OxmlElement(tag)
    flag.set(qn("w:val"), "true")
    properties.append(flag)


def shade(cell: Any, color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_margins(cell: Any, value: int = 100) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _highlight_segment(
    tag: str, side: str, old_tokens: list[str], new_tokens: list[str], span: tuple[int, ...]
) -> tuple[list[str], str | None] | None:
    """Tokens und Farbe eines Diff-Abschnitts für eine Seite; ``None`` = nicht zeigen."""
    i1, i2, j1, j2 = span
    if tag == "equal":
        return (old_tokens[i1:i2] if side == "old" else new_tokens[j1:j2]), None
    if side == "old" and tag in {"replace", "delete"}:
        return old_tokens[i1:i2], REMOVED
    if side == "new" and tag in {"replace", "insert"}:
        return new_tokens[j1:j2], ADDED
    return None


def add_word_highlight(paragraph: Any, old: str, new: str, side: str) -> None:
    from docx.shared import RGBColor

    old_tokens = (old or "").split()
    new_tokens = (new or "").split()
    matcher = difflib.SequenceMatcher(None, old_tokens, new_tokens)
    for tag, *span in matcher.get_opcodes():
        segment = _highlight_segment(tag, side, old_tokens, new_tokens, tuple(span))
        if segment is None or not segment[0]:
            continue
        tokens, color = segment
        if paragraph.runs:
            paragraph.add_run(" ")
        run = paragraph.add_run(" ".join(tokens))
        if color:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(color)
            if side == "old":
                run.font.strike = True


def add_labeled_value(paragraph: Any, label: str, value: str) -> None:
    if not value:
        return
    if paragraph.text:
        paragraph.add_run("\n")
    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    paragraph.add_run(value)


@dataclass(frozen=True)
class CellOptions:
    """Ausgabeschalter der Fassungszellen (aus den Ergebnis-Metadaten)."""

    include_answers: bool
    include_notes: bool
    highlight_words: bool


def fill_version_cell(cell: Any, row: CompareRow, side: str, options: CellOptions) -> None:
    old = side == "old"
    text = row.old_text if old else row.new_text
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    if not text:
        paragraph.add_run("nicht vorhanden" if old else "nicht mehr enthalten")
    elif options.highlight_words and row.status == "changed":
        add_word_highlight(paragraph, row.old_text, row.new_text, side)
    else:
        paragraph.add_run(text)
    if options.include_answers:
        add_labeled_value(paragraph, "Antwort", row.old_answer if old else row.new_answer)
        add_labeled_value(paragraph, "Bemerkung", row.old_comment if old else row.new_comment)
    if options.include_notes:
        add_labeled_value(paragraph, "Hinweis", row.old_note if old else row.new_note)
    if old:
        for run in paragraph.runs:
            run.italic = True


def reason_text(row: CompareRow) -> str:
    """Grundspalte: maschinelle Vorschläge gekennzeichnet, Prüfhinweis angehängt."""
    reason = row.reason
    if row.reason_source == "flowagent" and reason:
        reason = f"Maschineller Vorschlag (FlowAgent): {reason}"
    if row.reason_warning:
        reason = f"{reason}\nPrüfhinweis: {row.reason_warning}".strip()
    return reason
