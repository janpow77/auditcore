"""DOCX-Synopse (Extra ``docx-render``: python-docx).

Unveränderte Darstellung aus ``rendering.py`` des Originals: Querformat,
vierspaltige Tabelle (Fundstelle, bisherige/neue Fassung, Grund),
Gruppenzeilen, Wortmarkierung, Vermerk-Gliederung mit ``{{vergleich}}``,
nicht teilbare Tabellenzeilen, Seitenzahlen und konsolidierte Arbeitsfassung
bei der Gesetzessynopse. Zeitangaben werden in der lokalen Zeitzone des
Prozesses formatiert (wie im Original). Die Bibliothek liefert nur den
Renderer; ``auditcore_reporting`` 0.2.0 kennt keine DOCX-Ausgabe.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from auditcore_documents.docx_parts import (
    ADDED,
    HEADER_FILL,
    NAVY,
    REMOVED,
    TINT,
    CellOptions,
    DocxStyle,
    add_flag,
    fill_version_cell,
    format_timestamp,
    keep_row_together,
    reason_text,
    set_cell_margins,
    shade,
)
from auditcore_documents.errors import DependencyError
from auditcore_documents.model import ComparisonResult

__all__ = ["ADDED", "HEADER_FILL", "NAVY", "REMOVED", "TINT", "render_docx", "render_docx_bytes"]

#: Spaltenbreiten in cm: Fundstelle, bisherige Fassung, neue Fassung, Grund.
COLUMN_WIDTHS_CM = (4.7, 7.6, 7.6, 5.8)
DEFAULT_SECTIONS = ["changed", "removed", "added", "moved"]
DEFAULT_NOTICE = (
    "Die Maschine stellt Unterschiede fest und bereitet sie auf. "
    "Sie trifft keine Prüfungsentscheidung; die Würdigung bleibt beim Prüfer."
)


def _group_labels(mode: str) -> dict[str, str]:
    unit = "Textstellen" if mode == "text" else "Prüffragen"
    return {
        "changed": f"I. Geänderte {unit}",
        "removed": f"II. Entfallene {unit}",
        "added": f"III. Neue {unit}",
        # Gleicher Wortlaut, andere Stelle. Eigene Gruppe, damit eine
        # Umstellung nicht als Streichung samt Neuaufnahme gelesen wird.
        "moved": f"IV. Umgestellte {unit}",
        "unchanged": f"V. Unveränderte {unit}",
        "manual": "Von Hand zu bearbeiten",
    }


def _require_docx() -> Any:
    try:
        from docx import Document
    except ImportError as exc:
        raise DependencyError("Die DOCX-Ausgabe benötigt python-docx.") from exc
    return Document


def _setup_page(document: Any, style: DocxStyle) -> Any:
    """Querformat, Ränder und Grundschrift; liefert die nutzbare Breite."""
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, attr, Cm(style.margin_cm))
    normal = document.styles["Normal"]
    normal.font.name = style.font_family
    normal.font.size = Pt(style.body_size)
    normal.paragraph_format.line_spacing = style.line_spacing
    for attribute in ("w:ascii", "w:hAnsi", "w:cs"):
        normal._element.rPr.rFonts.set(qn(attribute), style.font_family)
    return section.page_width - section.left_margin - section.right_margin


def _add_header(section: Any, style: DocxStyle, user: str, usable_width: Any) -> None:
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Pt, RGBColor

    paragraph = section.header.paragraphs[0]
    paragraph.paragraph_format.tab_stops.clear_all()
    paragraph.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)
    lines = [line.strip() for line in style.header_text.splitlines() if line.strip()]
    texts = [f"{lines[0] if lines else 'Vermerk'}\t{user}", *(f"\n{line}" for line in lines[1:])]
    for text in texts:
        run = paragraph.add_run(text)
        run.font.size = Pt(9)
        run.font.name = style.font_family
        run.font.color.rgb = RGBColor.from_string(style.accent_color)


def _add_footer(section: Any, style: DocxStyle, usable_width: Any) -> None:
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.clear_all()
    footer.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)
    footer.add_run(style.footer_text)
    if not style.show_page_numbers:
        return
    for prefix, instruction in (("\tSeite ", "PAGE"), (" von ", "NUMPAGES")):
        footer.add_run(prefix)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), instruction)
        footer._p.append(field)


def _add_title_block(
    document: Any, result: ComparisonResult, title: str | None, style: DocxStyle
) -> None:
    from docx.shared import Pt, RGBColor

    title_run = document.add_paragraph().add_run(
        title or f"Vergleich: {result.old_filename} / {result.new_filename}"
    )
    title_run.bold = True
    title_run.font.size = Pt(style.heading_size)
    title_run.font.color.rgb = RGBColor.from_string(style.accent_color)
    metadata = result.metadata
    if style.show_file_metadata:
        old_modified = format_timestamp(metadata.get("old_modified_at"))
        new_modified = format_timestamp(metadata.get("new_modified_at"))
        document.add_paragraph(
            f"{metadata.get('old_label', 'Bisherige Fassung')}: "
            f"{result.old_filename} · geändert {old_modified} · SHA-256 {result.old_sha256}\n"
            f"{metadata.get('new_label', 'Neue Fassung')}: "
            f"{result.new_filename} · geändert {new_modified} · SHA-256 {result.new_sha256}"
        )
    mode_label = "Checkliste" if result.mode == "checklist" else "Fließtext"
    document.add_paragraph(
        f"Dokumentart: {mode_label} · Vergleichsmodul {result.version} · "
        f"Erstellt {format_timestamp(result.created_at)}"
    )
    notice_run = document.add_paragraph().add_run(
        str(metadata.get("work_aid_notice") or DEFAULT_NOTICE)
    )
    notice_run.italic = True


def _add_outline(document: Any, layout: dict[str, Any], style: DocxStyle) -> Any:
    """Vermerk-Gliederung; liefert den Absatz mit ``{{vergleich}}`` (Tabellenanker)."""
    from docx.shared import Pt, RGBColor

    anchor = None
    for raw_heading in layout.get("outline") or []:
        raw_heading = str(raw_heading).strip()
        if not raw_heading:
            continue
        is_comparison = "{{vergleich}}" in raw_heading
        heading_text = raw_heading.replace("{{vergleich}}", "").strip()
        heading_p = document.add_paragraph()
        heading_run = heading_p.add_run(heading_text or "Festgestellte Änderungen")
        heading_run.bold = True
        heading_run.font.size = Pt(max(style.body_size + 1, style.heading_size - 2))
        heading_run.font.color.rgb = RGBColor.from_string(style.accent_color)
        if is_comparison:
            anchor = heading_p
        else:
            document.add_paragraph()
    return anchor


def _add_table(document: Any, result: ComparisonResult, widths: list[Any]) -> Any:
    """Vierspaltige Tabelle mit wiederholter, nicht teilbarer Kopfzeile."""
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT

    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    metadata = result.metadata
    labels = [
        "Fundstelle",
        str(metadata.get("old_label") or "Bisherige Fassung"),
        str(metadata.get("new_label") or "Neue Fassung"),
        str(metadata.get("reason_label") or "Grund der Änderung"),
    ]
    for cell, label, width in zip(table.rows[0].cells, labels, widths, strict=True):
        cell.text = label
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, HEADER_FILL)
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    add_flag(table.rows[0]._tr.get_or_add_trPr(), "w:tblHeader")
    keep_row_together(table.rows[0])
    return table


def _add_group_row(table: Any, label: str) -> None:
    group_row = table.add_row()
    keep_row_together(group_row)
    cells = group_row.cells
    merged = cells[0].merge(cells[3])
    merged.text = label
    shade(merged, TINT)
    set_cell_margins(merged)
    for run in merged.paragraphs[0].runs:
        run.bold = True


def _add_data_row(table: Any, row: Any, widths: list[Any], options: CellOptions) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

    data_row = table.add_row()
    keep_row_together(data_row)
    cells = data_row.cells
    for cell, width in zip(cells, widths, strict=True):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
    cells[0].text = row.location or "—"
    cells[0].paragraphs[0].runs[0].bold = True
    fill_version_cell(cells[1], row, "old", options)
    fill_version_cell(cells[2], row, "new", options)
    cells[3].text = reason_text(row)


def _add_rows(table: Any, result: ComparisonResult, widths: list[Any]) -> None:
    """Ausgewählte Zeilen der gewählten Abschnitte, je Statuswechsel eine Gruppenzeile."""
    metadata = result.metadata
    group_labels = _group_labels(result.mode)
    output_sections = set(metadata.get("output_sections") or DEFAULT_SECTIONS)
    options = CellOptions(
        include_answers=bool(metadata.get("include_answers", True)),
        include_notes=bool(metadata.get("include_notes", True)),
        highlight_words=bool(metadata.get("highlight_words", True)),
    )
    current_group = ""
    for row in result.rows:
        if not row.selected or row.status not in output_sections:
            continue
        if row.status != current_group:
            current_group = row.status
            _add_group_row(table, group_labels.get(row.status, row.status))
        _add_data_row(table, row, widths, options)


def _add_open_commands(document: Any, result: ComparisonResult, style: DocxStyle) -> None:
    from docx.shared import RGBColor

    open_commands = result.metadata.get("open_commands") or []
    if not open_commands:
        return
    heading_run = document.add_paragraph().add_run("Von Hand zu bearbeiten")
    heading_run.bold = True
    heading_run.font.color.rgb = RGBColor.from_string(style.accent_color)
    for command in open_commands:
        document.add_paragraph(str(command), style="List Bullet")


def _add_article_law_appendix(document: Any, result: ComparisonResult) -> None:
    """Befehlszählung und konsolidierte Arbeitsfassung der Gesetzessynopse."""
    from docx.shared import Pt

    metadata = result.metadata
    if metadata.get("comparison_type") != "article_law":
        return
    document.add_paragraph(
        f"Änderungsbefehle erkannt: {metadata.get('recognised_commands', 0)} · "
        f"offen: {metadata.get('open_command_count', 0)}"
    )
    consolidated = metadata.get("consolidated_text") or []
    if not consolidated or not metadata.get("include_consolidated_text", True):
        return
    document.add_page_break()
    run = document.add_paragraph().add_run("Konsolidierte Arbeitsfassung (nicht amtlich)")
    run.bold = True
    run.font.size = Pt(13)
    last_section = ""
    for item in consolidated:
        if item.get("repealed"):
            continue
        if item.get("section") != last_section:
            last_section = str(item.get("section") or "")
            document.add_paragraph().add_run(last_section).bold = True
        document.add_paragraph(str(item.get("text") or ""))


def _build_document(
    result: ComparisonResult,
    *,
    title: str | None = None,
    user: str = "",
    header: str | None = None,
    profile: str = "memo",
    layout: dict[str, Any] | None = None,
) -> Any:
    document_class = _require_docx()
    from docx.shared import Cm

    layout = layout or {}
    style = DocxStyle.from_layout(layout, header)
    document = document_class()
    usable_width = _setup_page(document, style)
    section = document.sections[0]
    _add_header(section, style, user, usable_width)
    _add_footer(section, style, usable_width)
    _add_title_block(document, result, title, style)
    anchor = _add_outline(document, layout, style) if profile == "memo" else None
    widths = [Cm(width) for width in COLUMN_WIDTHS_CM]
    table = _add_table(document, result, widths)
    _add_rows(table, result, widths)
    if anchor is not None:
        anchor._p.addnext(table._tbl)
    _add_open_commands(document, result, style)
    _add_article_law_appendix(document, result)
    return document


def render_docx(
    result: ComparisonResult,
    output: Path,
    *,
    title: str | None = None,
    user: str = "",
    header: str | None = None,
    profile: str = "memo",
    layout: dict[str, Any] | None = None,
) -> None:
    """Synopse als DOCX nach ``output`` schreiben (Verzeichnisse werden angelegt).

    Ungültige Layoutwerte (z. B. nicht numerische Schriftgröße) lösen wie im
    Original ``ValueError`` aus; fehlt python-docx, ``DependencyError``.
    """
    document = _build_document(
        result, title=title, user=user, header=header, profile=profile, layout=layout
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def render_docx_bytes(
    result: ComparisonResult,
    *,
    title: str | None = None,
    user: str = "",
    header: str | None = None,
    profile: str = "memo",
    layout: dict[str, Any] | None = None,
) -> bytes:
    """Wie :func:`render_docx`, aber als Bytes ohne Dateizugriff."""
    import io

    document = _build_document(
        result, title=title, user=user, header=header, profile=profile, layout=layout
    )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
