"""Table readers for transparency lists (CSV with the standard library, XLSX optional).

Header detection follows ``flowworkshop/services/dataframe_service.py``
(``_read_csv_smart``/``_read_excel_smart``) without pandas. Two explicit
typing modes:

* ``"legacy"`` reproduces the pandas type inference of the source application
  (numeric columns become ``int``/``float``, gaps turn integer columns into
  floats, the default NA strings become missing). It keeps the record hashes
  of existing inventories stable (FS-W02/W03).
* ``"text"`` keeps every cell as the original text (``None`` only for empty
  cells), so ``"01067"`` stays ``"01067"``. It changes identities of rows whose
  legacy value was re-typed and therefore needs a decision of the consumer.

Resource limits apply to every read: file size, rows, columns and cell length.
"""

from __future__ import annotations

import csv
import io
import math
import re
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from typing import Any, Literal

from .errors import OptionalDependencyError, SourceFormatError

Typing = Literal["legacy", "text"]

#: pandas ``read_csv`` default NA strings (``keep_default_na=True``).
PANDAS_NA = frozenset(
    {
        "",
        "#N/A",
        "#N/A N/A",
        "#NA",
        "-1.#IND",
        "-1.#QNAN",
        "-NaN",
        "-nan",
        "1.#IND",
        "1.#QNAN",
        "<NA>",
        "N/A",
        "NA",
        "NULL",
        "NaN",
        "None",
        "n/a",
        "nan",
        "null",
    }
)
_INT_RE = re.compile(r"[+-]?\d+")
_FLOAT_RE = re.compile(
    r"[+-]?(?:\d+\.?\d*(?:[eE][+-]?\d+)?|\.\d+(?:[eE][+-]?\d+)?|inf|infinity|nan)", re.IGNORECASE
)


@dataclass(frozen=True)
class Limits:
    """Upper bounds for one source file."""

    max_bytes: int = 64 * 1024 * 1024
    max_rows: int = 500_000
    max_columns: int = 500
    max_cell_characters: int = 100_000


@dataclass(frozen=True)
class Table:
    """Header names and data rows of one sheet, plus how they were detected."""

    headers: tuple[str, ...]
    rows: tuple[tuple[Any, ...], ...]
    header_row: int
    delimiter: str | None
    encoding: str | None
    typing: str


def _check_size(content: bytes, limits: Limits) -> None:
    if not isinstance(content, (bytes, bytearray)):
        raise SourceFormatError("Der Dateiinhalt ist als Bytes zu übergeben.")
    if len(content) > limits.max_bytes:
        raise SourceFormatError("Die Datei überschreitet die zulässige Größe.")


def _unique_headers(raw: Sequence[object]) -> tuple[str, ...]:
    """pandas header mangling: empty → ``Unnamed: i``, duplicates → ``name.1``."""
    names: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(raw):
        name = "" if value is None else str(value)
        if name == "" or (isinstance(value, float) and math.isnan(value)):
            name = f"Unnamed: {index}"
        base = name
        count = seen.get(base, 0)
        while name in seen:
            count += 1
            name = f"{base}.{count}"
        seen[base] = count
        seen[name] = 0
        names.append(name)
    return tuple(names)


def _infer_column(values: list[Any]) -> list[Any]:
    """pandas-like inference for one column of text cells."""
    cleaned = [None if v is None or v in PANDAS_NA else v for v in values]
    present = [v for v in cleaned if v is not None]
    if not present:
        return [math.nan] * len(values)
    stripped = [str(v).strip() for v in present]
    if all(_INT_RE.fullmatch(s) for s in stripped):
        if len(present) == len(cleaned):
            return [int(str(v).strip()) for v in cleaned]
        return [math.nan if v is None else float(str(v).strip()) for v in cleaned]
    if all(_FLOAT_RE.fullmatch(s) for s in stripped):
        return [math.nan if v is None else float(str(v).strip()) for v in cleaned]
    lowered = {s.lower() for s in stripped}
    if lowered <= {"true", "false"} and len(present) == len(cleaned):
        return [str(v).strip().lower() == "true" for v in cleaned]
    return [math.nan if v is None else v for v in cleaned]


def _typed(rows: list[list[Any]], width: int, typing: Typing) -> list[tuple[Any, ...]]:
    padded = [row + [None] * (width - len(row)) for row in rows]
    if typing == "text":
        return [tuple(None if (c is None or c == "") else c for c in row) for row in padded]
    columns = [_infer_column([row[i] for row in padded]) for i in range(width)]
    return [tuple(column[r] for column in columns) for r in range(len(padded))]


def _decode(content: bytes) -> tuple[str, str]:
    for encoding in ("utf-8", "cp1252", "latin-1"):
        try:
            return content.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace"), "utf-8-replace"  # pragma: no cover


HeaderDetection = Literal["legacy", "strict"]


def read_csv(
    content: bytes,
    *,
    typing: Typing = "legacy",
    header_detection: HeaderDetection = "legacy",
    limits: Limits | None = None,
) -> Table:
    """Transparency-list CSV: encoding and delimiter detection, title rows, double headers.

    Unlike the source application, title rows with fewer fields than the
    header do not abort the file (FS-W04). ``header_detection="legacy"`` keeps
    the source rule for a second, machine-readable header row, which also
    takes a first *data* row with three or more purely numeric cells as header
    and thereby loses it (FS-W08); ``"strict"`` requires such header cells to
    contain a letter.
    """
    limits = limits or Limits()
    _check_size(content, limits)
    text, encoding = _decode(content)
    delimiter, records = _csv_records(text, limits)
    header = _pick_header(
        [[v.strip() for v in r if v.strip()] for r in records[:25]], min_ratio=0.6
    )
    header = _second_header(records, header, header_detection)
    headers = _unique_headers(records[header])
    if len(headers) > limits.max_columns:
        raise SourceFormatError("Die Datei enthält mehr Spalten als zulässig.")
    data = records[header + 1 :]
    _check_csv_rows(data, len(headers), limits)
    return Table(
        headers,
        tuple(_typed([list(r) for r in data], len(headers), typing)),
        header,
        delimiter,
        encoding,
        typing,
    )


def _csv_records(text: str, limits: Limits) -> tuple[str, list[list[str]]]:
    """Sniffed delimiter (``;`` as fallback) and non-blank records."""
    try:
        delimiter = csv.Sniffer().sniff(text[:8192], delimiters=";,\t|").delimiter
    except csv.Error:
        delimiter = ";"
    try:
        records = [r for r in csv.reader(io.StringIO(text), delimiter=delimiter) if r]
    except csv.Error as exc:
        raise SourceFormatError(f"CSV nicht lesbar: {exc}") from exc
    if not records:
        raise SourceFormatError("Die CSV-Datei ist leer.")
    if len(records) > limits.max_rows + 30:
        raise SourceFormatError("Die Datei enthält mehr Zeilen als zulässig.")
    return delimiter, records


def _pick_header(non_empty_rows: Sequence[Sequence[str]], *, min_ratio: float) -> int:
    """Source heuristic: first row with ≥ 3 cells and enough text, else the most textual."""
    candidates: list[tuple[int, int, float]] = []
    for idx, cells in enumerate(non_empty_rows):
        if len(cells) < 2:
            continue
        ratio = sum(1 for c in cells if any(ch.isalpha() for ch in c)) / max(len(cells), 1)
        candidates.append((idx, len(cells), ratio))
    for idx, count, ratio in candidates:
        if count >= 3 and ratio >= min_ratio:
            return idx
    if candidates:
        return max(candidates, key=lambda item: (item[2], item[1]))[0]
    return 0


def _second_header(records: Sequence[Sequence[str]], header: int, mode: HeaderDetection) -> int:
    """Skip a machine-readable second header (``snake_case``) below the header row."""
    if header + 1 >= len(records):
        return header
    following = [v.strip() for v in records[header + 1] if v.strip()]
    snake = sum(
        1
        for v in following
        if re.fullmatch(r"[a-z0-9_]+", v) and (mode == "legacy" or re.search(r"[a-z]", v))
    )
    if following and snake >= max(3, len(following) // 2):
        return header + 1
    return header


def _check_csv_rows(data: Sequence[Sequence[str]], width: int, limits: Limits) -> None:
    for row in data:
        if len(row) > width:
            raise SourceFormatError(
                f"Zeile mit {len(row)} Feldern bei {width} Spalten; Datei nicht eindeutig lesbar."
            )
        if any(len(cell) > limits.max_cell_characters for cell in row):
            raise SourceFormatError("Eine Zelle überschreitet die zulässige Länge.")


def _openpyxl() -> Any:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - exercised without the extra
        raise OptionalDependencyError(
            "XLSX-Dateien benötigen das Extra auditcore_funding_sources[xlsx] (openpyxl)."
        ) from exc
    return openpyxl


def _excel_value(value: object, typing: Typing) -> object:
    if typing == "text":
        if value is None:
            return None
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        if isinstance(value, (datetime, date, time)):
            return value.isoformat()
        text = str(value)
        return text if text != "" else None
    return value


def _numeric_text(value: object) -> bool:
    return isinstance(value, str) and bool(
        _INT_RE.fullmatch(value.strip()) or _FLOAT_RE.fullmatch(value.strip())
    )


def _infer_excel_column(values: list[Any]) -> list[Any]:
    """pandas ``read_excel`` rules: all-numeric columns (numbers or numeric text) become
    ``int`` without gaps or ``float`` with gaps; other columns keep their cell values."""
    missing = [v is None or (isinstance(v, str) and v in PANDAS_NA) for v in values]
    present = [v for v, gap in zip(values, missing, strict=True) if not gap]
    if not present:
        return [math.nan] * len(values)
    if all(isinstance(v, bool) for v in present) and not any(missing):
        return values
    numeric = all(
        (isinstance(v, (int, float)) and not isinstance(v, bool)) or _numeric_text(v)
        for v in present
    )
    if not numeric:
        return [math.nan if gap else v for v, gap in zip(values, missing, strict=True)]
    integral = all(
        (isinstance(v, int) and not isinstance(v, bool))
        or (isinstance(v, str) and _INT_RE.fullmatch(v.strip()))
        for v in present
    )
    if integral and not any(missing):
        return [int(v) if isinstance(v, str) else v for v in values]
    return [math.nan if gap else float(v) for v, gap in zip(values, missing, strict=True)]


def read_xlsx(
    content: bytes,
    *,
    sheet: str | int | None = 0,
    header_row: int = 0,
    typing: Typing = "legacy",
    limits: Limits | None = None,
) -> Table:
    """First/selected sheet; header detection as in the source unless ``header_row`` > 0."""
    limits = limits or Limits()
    _check_size(content, limits)
    rows = _xlsx_rows(content, sheet, limits)
    header = header_row
    if not header_row or header_row <= 0:
        header = _pick_header(
            [[str(c).strip() for c in r if c is not None and str(c).strip()] for r in rows[:31]],
            min_ratio=0.4,
        )
    if header >= len(rows):
        raise SourceFormatError("Die angegebene Kopfzeile liegt hinter dem Tabellenende.")
    width = _used_width(rows, header)
    headers = _unique_headers(rows[header][:width])
    data = [(r + [None] * width)[:width] for r in rows[header + 1 :]]
    if typing == "text":
        typed = [tuple(_excel_value(v, typing) for v in r) for r in data]
    else:
        columns = [_infer_excel_column([r[i] for r in data]) for i in range(width)]
        typed = [tuple(c[j] for c in columns) for j in range(len(data))]
    return Table(headers, tuple(typed), header, None, None, typing)


def _xlsx_rows(content: bytes, sheet: str | int | None, limits: Limits) -> list[list[Any]]:
    """Cell values of the selected sheet without trailing empty rows."""
    openpyxl = _openpyxl()
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any parser failure is a format error
        raise SourceFormatError(f"XLSX nicht lesbar: {type(exc).__name__}") from exc
    try:
        if isinstance(sheet, str) and sheet in workbook.sheetnames:
            ws = workbook[sheet]
        elif isinstance(sheet, int) and sheet < len(workbook.worksheets):
            ws = workbook.worksheets[sheet]
        else:
            ws = workbook.worksheets[0]
        rows: list[list[Any]] = []
        for index, row in enumerate(ws.iter_rows(values_only=True)):
            if index > limits.max_rows + 30:
                raise SourceFormatError("Die Datei enthält mehr Zeilen als zulässig.")
            if len(row) > limits.max_columns:
                raise SourceFormatError("Die Datei enthält mehr Spalten als zulässig.")
            rows.append(list(row))
    finally:
        workbook.close()
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    return rows


def _used_width(rows: Sequence[Sequence[object]], header: int) -> int:
    """Header width without trailing columns that are empty in header and data."""
    raw = rows[header]
    width = len(raw)
    while (
        width
        and raw[width - 1] is None
        and all((len(r) < width or r[width - 1] is None) for r in rows[header + 1 :])
    ):
        width -= 1
    return width


def read_table(
    content: bytes,
    file_name: str,
    *,
    sheet: str | int | None = None,
    header_row: int = 0,
    typing: Typing = "legacy",
    header_detection: HeaderDetection = "legacy",
    limits: Limits | None = None,
) -> Table:
    """Dispatch by file extension (``csv``, ``xlsx``/``xlsm``); PDF stays with the consumer."""
    ext = file_name.rsplit(".", 1)[-1].lower() if "." in (file_name or "") else ""
    if ext == "csv":
        return read_csv(content, typing=typing, header_detection=header_detection, limits=limits)
    if ext in ("xlsx", "xlsm"):
        return read_xlsx(
            content,
            sheet=sheet if sheet is not None else 0,
            header_row=header_row,
            typing=typing,
            limits=limits,
        )
    raise SourceFormatError(
        f"Nur CSV und XLSX werden von dieser Bibliothek gelesen, nicht '{ext}' "
        f"(file_name={file_name})."
    )
