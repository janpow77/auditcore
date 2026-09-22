"""Observe the exact MIT Flowlib modules before adapting workbook export."""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import io
import json
import sys
import types
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

HASHES = {
    "formats": "f86f2e1f01397e0a886568789713705673b886e9dfd11212e12286daf3154c35",
    "styles": "dae32824361c9b61b87f5dd18e719439030918ed6239867f5531ecb850928aaa",
    "report": "f28f9e740bfd19e75bace228c0a8d840c4f7f63d6db5a73cf3664894d84242bd",
}


def snapshot(ws):
    """Capture observable cell values, selected styles and widths after XLSX roundtrip."""
    return {
        "cells": [
            {
                "coordinate": c.coordinate,
                "value": c.value.isoformat() if hasattr(c.value, "isoformat") else c.value,
                "type": c.data_type,
                "number_format": c.number_format,
                "bold": c.font.bold,
                "font_name": c.font.name,
                "fill": c.fill.fgColor.rgb if c.fill.fill_type else None,
                "border": c.border.left.style,
                "horizontal": c.alignment.horizontal,
            }
            for row in ws
            for c in row
        ],
        "widths": {key: dimension.width for key, dimension in ws.column_dimensions.items()},
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    # Load the exact source modules, avoiding unrelated application/package __init__ imports.
    for name in ("flowlib", "flowlib.excel"):
        module = types.ModuleType(name)
        module.__path__ = []
        sys.modules[name] = module
    for name, expected in HASHES.items():
        path = args.source / f"{name}.py"
        assert hashlib.sha256(path.read_bytes()).hexdigest() == expected
        spec = importlib.util.spec_from_file_location(f"flowlib.excel.{name}", path)
        module = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = module
        spec.loader.exec_module(module)
    report = sys.modules["flowlib.excel.report"]
    records = []
    cases = [
        (
            "ordinary",
            ["Betrag", "Quote", "Datum", "Text"],
            [[123.45, 0.25, date(2024, 1, 2), "Änderung"], [None, 0.1, None, "longer sample text"]],
            1,
            True,
            True,
        ),
        (
            "offset_no_style",
            ["Anzahl", "Stunden", "Name"],
            [[7, 2.5, "A"], [8, 3.0, "B"]],
            3,
            False,
            False,
        ),
        ("empty", ["Betrag", "Text"], [], 1, True, True),
        ("formula_legacy", ["Text"], [["=1+1"], ["+cmd"], ["@SUM(A1)"]], 1, True, True),
        ("nan", ["Betrag"], [[float("nan")], [10.0]], 1, True, True),
    ]
    for name, columns, rows, start, zebra, border in cases:
        wb = Workbook()
        ws = wb.active
        next_row = report.write_dataframe(
            ws, pd.DataFrame(rows, columns=columns), start_row=start, zebra=zebra, border=border
        )
        stream = io.BytesIO()
        wb.save(stream)
        observed = snapshot(load_workbook(io.BytesIO(stream.getvalue())).active)
        records.append(
            {
                "name": name,
                "columns": columns,
                "rows": [
                    [
                        v.isoformat()
                        if isinstance(v, date)
                        else None
                        if isinstance(v, float) and pd.isna(v)
                        else v
                        for v in row
                    ]
                    for row in rows
                ],
                "start_row": start,
                "zebra": zebra,
                "border": border,
                "next_row": next_row,
                "observed": observed,
            }
        )
    args.output.write_text(
        json.dumps(
            {
                "source_commit": "aca2dc6aad25aea0720312dbcc6da00b0bcba330",
                "source_hashes": HASHES,
                "license": "MIT",
                "records": records,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n"
    )
    print(f"Observed {len(records)} actual Flowlib workbook cases")


if __name__ == "__main__":
    main()
