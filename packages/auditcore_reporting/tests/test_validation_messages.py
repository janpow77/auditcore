"""Pin each table validation branch, its exception type and message order."""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from auditcore_reporting import ReportTable, WorkbookLimitError, WorkbookLimits, render_workbook
from auditcore_reporting.workbook import ExcelOptions

CASES = [
    (ReportTable("", ["A"], []), ValueError, "1 to 31"),
    (ReportTable("x" * 32, ["A"], []), ValueError, "1 to 31"),
    (ReportTable("a/b", ["A"], []), ValueError, "Invalid worksheet name"),
    (ReportTable("'a", ["A"], []), ValueError, "Invalid worksheet name"),
    (ReportTable("a\x01", ["A"], []), ValueError, "XML 1.0"),
    (ReportTable("S", "AB", []), TypeError, "sequence of names"),
    (ReportTable("S", [], []), WorkbookLimitError, "excessive column"),
    (ReportTable("S", ["A", ""], []), ValueError, "nonempty strings"),
    (ReportTable("S", ["A", "A"], []), ValueError, "Duplicate column"),
    (ReportTable("S", ["A"], [], start_row=True), TypeError, "start_row must"),
    (ReportTable("S", ["A"], [], start_row=0), WorkbookLimitError, "start_row exceeds"),
    (ReportTable("S", ["A"], [], profile="x"), ValueError, "Unknown format profile"),
    (ReportTable("S", ["A"], [], formats={"B": "0"}), ValueError, "unknown column"),
    (ReportTable("S", ["A"], [], formats={"A": ""}), ValueError, "1 to 255"),
    # Several defects at once: the first check in the historical order wins.
    (ReportTable("a/b", ["A", "A"], [], start_row=0), ValueError, "Invalid worksheet name"),
    (ReportTable("S", ["A", "A"], [], start_row=0), ValueError, "Duplicate column"),
    (ReportTable("S", ["A"], [], start_row=0, profile="x"), WorkbookLimitError, "start_row"),
]


@pytest.mark.parametrize(("table", "error", "message"), CASES)
def test_table_validation_branches_keep_type_message_and_order(table, error, message):
    with pytest.raises(error, match=message):
        render_workbook([table])


def test_column_limit_uses_configured_maximum():
    options = ExcelOptions(limits=WorkbookLimits(max_columns=1))
    with pytest.raises(WorkbookLimitError, match="excessive column"):
        render_workbook([ReportTable("S", ["A", "B"], [])], options)


def test_sheet_finish_sets_widths_freeze_and_filter_over_all_columns():
    rows = [["x" * 70, 1], ["y", None]]
    output = render_workbook([ReportTable("S", ["Name", "Anzahl"], rows, start_row=3)])
    ws = load_workbook(io.BytesIO(output)).active
    assert ws.freeze_panes == "A4"
    assert ws.auto_filter.ref == "A3:B5"
    assert ws.column_dimensions["A"].width == 50
    assert ws.column_dimensions["B"].width == 8
    assert ws["B4"].number_format == "#,##0"
    assert ws["A5"].fill.fgColor.rgb.endswith("F2F2F2")
