"""Page orientation, table layout and TrueType fonts – issue #235 items 2–4."""

from __future__ import annotations

import io
import zipfile
from dataclasses import replace
from pathlib import Path
from typing import Any

import pytest
import reportlab
from docx import Document
from pypdf import PdfReader
from reportlab.lib.colors import HexColor
from reportlab.lib.rl_accel import fp_str

from auditcore_reporting.templates import (
    NEUTRAL_DESIGN,
    DesignProfile,
    PdfFont,
    TemplateError,
    builtin_registry,
    define_template,
    design_from_dict,
    render,
    resolve,
)
from auditcore_reporting.templates.resolve import RTable

SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {
        "titel": {"type": "string"},
        "posten": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "name": {"type": "string"},
                    "betrag": {"type": "number"},
                    "strittig": {"type": "boolean"},
                },
            },
        },
    },
}
DATA = {
    "titel": "Preisvergleich",
    "posten": [
        {"name": "A", "betrag": 10, "strittig": False},
        {"name": "B", "betrag": 2000, "strittig": True},
        {"name": "C", "betrag": 30, "strittig": False},
    ],
}
COLUMNS = [
    {"header": "Name", "cell": "{{ p.name }}", "width": 3},
    {"header": "Betrag", "cell": "{{ p.betrag | eur }}", "align": "right", "width": 1,
     "fill": [{"if": "p.strittig", "color": "F8CBAD"}]},
]  # fmt: skip
STYLED = {
    "type": "table", "source": "posten", "as": "p", "columns": COLUMNS, "borders": "horizontal",
    "header_fill": "FFF2CC", "stripe": "F2F2F2",
    "row_fill": [{"if": {"greater": ["p.betrag", 1000]}, "color": "DDEBF7", "bold": True}],
}  # fmt: skip
VERA = Path(reportlab.__file__).parent / "fonts"
DEJAVU = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")


def template(*blocks: dict[str, Any], **extra: object) -> Any:
    return define_template(
        {"id": "layout", "version": "1.0.0", "title": "{{ titel }}", "schema": SCHEMA,
         "blocks": list(blocks) or [{"type": "paragraph", "text": "{{ titel }}"}], **extra}
    )  # fmt: skip


def docx_xml(content: bytes, part: str = "word/document.xml") -> str:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return archive.read(part).decode()


# --- 3. Seitenausrichtung ------------------------------------------------------------------

LANDSCAPE = replace(NEUTRAL_DESIGN, id="quer", orientation="landscape")


@pytest.mark.parametrize(
    ("extra", "design"),
    [({"orientation": "landscape"}, NEUTRAL_DESIGN), ({}, LANDSCAPE)],
    ids=["vorlage", "gestaltung"],
)
def test_landscape_in_every_format(extra: dict[str, object], design: DesignProfile) -> None:
    quer = template(**extra)
    docx = render(quer, DATA, "docx", design).content
    assert '<w:pgSz w:w="16838" w:h="11906" w:orient="landscape"/>' in docx_xml(docx)
    section = Document(io.BytesIO(docx)).sections[0]
    assert section.page_width > section.page_height
    box = PdfReader(io.BytesIO(render(quer, DATA, "pdf", design).content)).pages[0].mediabox
    assert float(box.width) > float(box.height)
    assert "@page{size:A4 landscape}" in render(quer, DATA, "html", design).content.decode()


def test_template_orientation_overrides_the_design() -> None:
    hoch = template(orientation="portrait")
    docx = render(hoch, DATA, "docx", LANDSCAPE).content
    assert '<w:pgSz w:w="11906" w:h="16838"/>' in docx_xml(docx)


def test_orientation_is_validated() -> None:
    with pytest.raises(TemplateError, match="orientation"):
        template(orientation="quer")
    with pytest.raises(TemplateError, match="orientation"):
        design_from_dict({"orientation": "quer"})
    assert design_from_dict({"id": "quer", "orientation": "landscape"}).orientation == "landscape"


# --- 4. Tabellen ---------------------------------------------------------------------------


def test_styled_table_resolves_widths_fills_and_emphasis() -> None:
    (node,) = resolve(template(STYLED), DATA).nodes
    assert isinstance(node, RTable)
    assert node.widths == (3.0, 1.0) and node.borders == "horizontal"
    assert node.header_fill == "FFF2CC"
    # Row 2 matches the row rule (bold, blue); its amount cell has its own colour.
    assert node.fills == (("", ""), ("DDEBF7", "F8CBAD"), ("", ""))
    assert node.bold == ((False, False), (True, True), (False, False))


def test_stripe_and_column_emphasis() -> None:
    columns = [{**COLUMNS[0], "bold": True}, {"header": "Betrag", "cell": "x"}]
    table = {"type": "table", "source": "posten", "as": "p", "columns": columns, "stripe": "F2F2F2"}
    (node,) = resolve(template(table), DATA).nodes
    assert isinstance(node, RTable)
    assert node.fills == (("", ""), ("F2F2F2", "F2F2F2"), ("", ""))
    assert node.bold == ((True, False), (True, False), (True, False))


def test_plain_table_has_no_layout_extras() -> None:
    plain = {"type": "table", "source": "posten", "as": "p", "columns": COLUMNS[:1]}
    (node,) = resolve(template(plain), DATA).nodes
    assert node == RTable(("Name",), ("left",), (("A",), ("B",), ("C",)), (3.0,))


def test_styled_table_in_docx() -> None:
    content = render(template(STYLED), DATA, "docx").content
    xml = docx_xml(content)
    assert '<w:gridCol w:w="6802"/><w:gridCol w:w="2267"/>' in xml
    assert "TableGrid" not in xml and "<w:insideH " in xml and "<w:insideV " not in xml
    assert 'w:fill="FFF2CC"' in xml and 'w:fill="F8CBAD"' in xml and 'w:fill="DDEBF7"' in xml
    table = Document(io.BytesIO(content)).tables[0]
    assert [c.text for c in table.rows[2].cells] == ["B", "2.000,00 €"]
    assert all(run.bold for cell in table.rows[2].cells for run in cell.paragraphs[0].runs)
    assert not any(run.bold for run in table.rows[1].cells[0].paragraphs[0].runs)


def _pdf_fill(hex_color: str) -> str:
    color = HexColor(f"#{hex_color}")
    return f"{fp_str(color.red, color.green, color.blue)} rg"


def test_styled_table_in_pdf() -> None:
    page = PdfReader(io.BytesIO(render(template(STYLED), DATA, "pdf").content)).pages[0]
    stream = page.get_contents().get_data().decode("latin-1")  # type: ignore[union-attr]
    for color in ("FFF2CC", "F8CBAD", "DDEBF7"):
        assert _pdf_fill(color) in stream, color
    assert "2.000,00 €" in page.extract_text()


def test_styled_table_in_html() -> None:
    html = render(template(STYLED), DATA, "html").content.decode()
    assert '<col style="width:75.00%"><col style="width:25.00%">' in html
    assert "background:#FFF2CC" in html and "background:#F8CBAD" in html
    assert "font-weight:bold" in html and "border-left:none" in html


@pytest.mark.parametrize(
    ("change", "message"),
    [
        ({"borders": "dick"}, "borders"),
        ({"stripe": "f2f2f2"}, "RRGGBB"),
        ({"row_fill": [{"color": "FFFFFF", "farbe": 1}]}, "unbekannte"),
        ({"row_fill": [{"if": "unbekannt", "color": "FFFFFF"}]}, "nicht deklariert"),
        ({"columns": [{"header": "x", "cell": "y", "width": -1}]}, "width"),
        ({"columns": [{"header": "x", "cell": "y", "bold": "ja"}]}, "bold"),
    ],
    ids=["borders", "farbe", "unbekannt", "bedingung", "breite", "fett"],
)
def test_table_layout_is_validated(change: dict[str, object], message: str) -> None:
    with pytest.raises(TemplateError, match=message):
        template({**STYLED, **change})


# --- 2. Schriften --------------------------------------------------------------------------


def vera_design(**change: object) -> DesignProfile:
    font = PdfFont(
        "Vera", (VERA / "Vera.ttf").read_bytes(), bold=(VERA / "VeraBd.ttf").read_bytes()
    )
    return DesignProfile(id="vera", pdf_font="Vera", pdf_fonts=(font,), **change)  # type: ignore[arg-type]


def test_true_type_font_carries_characters_beyond_win_ansi() -> None:
    data = {**DATA, "titel": "Ω-Prüfung"}
    titled = template(
        {"type": "heading", "text": "{{ titel }}"}, {"type": "paragraph", "text": "{{ titel }}"}
    )
    base = render(titled, data, "pdf").content
    assert "Ω" not in PdfReader(io.BytesIO(base)).pages[0].extract_text()
    content = render(titled, data, "pdf", vera_design()).content
    assert "Ω-Prüfung" in PdfReader(io.BytesIO(content)).pages[0].extract_text()
    assert b"BitstreamVeraSans-Roman" in content and b"BitstreamVeraSans-Bold" in content
    assert content == render(titled, data, "pdf", vera_design()).content


@pytest.mark.skipif(not DEJAVU.is_file(), reason="DejaVu Sans nicht installiert")
def test_check_boxes_with_a_system_font() -> None:
    design = DesignProfile(
        id="dejavu", pdf_font="DejaVu", pdf_fonts=(PdfFont("DejaVu", DEJAVU.read_bytes()),)
    )
    content = render(template(), {**DATA, "titel": "☐ offen ☒ erledigt"}, "pdf", design).content
    assert "☐ offen ☒ erledigt" in PdfReader(io.BytesIO(content)).pages[0].extract_text()


def test_font_profile_is_validated_and_serialisable() -> None:
    design = vera_design()
    fonts = design.to_dict()["pdf_fonts"]
    assert isinstance(fonts, list) and fonts[0]["name"] == "Vera" and fonts[0]["bold"] is True
    assert len(fonts[0]["sha256"]) == 64
    with pytest.raises(TemplateError, match="pdf_font"):
        DesignProfile(id="x-y", pdf_font="Vera")
    with pytest.raises(TemplateError, match="TrueType"):
        PdfFont("Kaputt", b"not a font")
    with pytest.raises(TemplateError, match="Name"):
        PdfFont("Helvetica", (VERA / "Vera.ttf").read_bytes())
    with pytest.raises(TemplateError, match="Python-API"):
        design_from_dict({"pdf_fonts": fonts})
    assert design_from_dict({"pdf_fonts": []}) == NEUTRAL_DESIGN


def test_unreadable_true_type_file_is_reported() -> None:
    broken = PdfFont("Kaputt", b"\x00\x01\x00\x00" + b"\x00" * 64)
    design = DesignProfile(id="kaputt", pdf_font="Kaputt", pdf_fonts=(broken,))
    with pytest.raises(TemplateError, match="nicht lesbar"):
        render(template(), DATA, "pdf", design)


def test_true_type_font_does_not_change_docx_or_html() -> None:
    memo = builtin_registry().get("vermerk")
    for output in ("docx", "html"):
        plain = render(memo, memo.sample, output, replace(vera_design(), pdf_font="Helvetica"))
        assert plain.content == render(memo, memo.sample, output, vera_design()).content


@pytest.mark.parametrize("output", ["docx", "pdf", "html"])
def test_builtin_templates_unchanged_by_new_defaults(output: str) -> None:
    for name in ("vermerk", "pruefbericht"):
        built = builtin_registry().get(name)
        first = render(built, built.sample, output).content
        assert first == render(built, built.sample, output, NEUTRAL_DESIGN).content
        if output == "docx":
            assert "w:orient" not in docx_xml(first)
