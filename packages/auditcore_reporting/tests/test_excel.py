"""Real workbook output, observed legacy styling and deliberate security corrections."""

from __future__ import annotations

import copy
import io
import json
from dataclasses import replace
from datetime import UTC, date, datetime
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from auditcore_reporting import (
    ExcelOptions,
    ReportTable,
    WorkbookLimitError,
    WorkbookLimits,
    get_number_format,
    get_profile_format,
    render_workbook,
)

GOLDEN = json.loads((Path(__file__).parent / "fixtures/flowlib-workbook-golden.json").read_text())


def snapshot(ws):
    return {
        "cells": [
            {
                "coordinate": c.coordinate,
                "value": c.value.isoformat() if hasattr(c.value, "isoformat") else c.value,
                "type": c.data_type,
                "number_format": c.number_format,
                "bold": c.font.bold,
                "font_name": c.font.name,
                "fill": c.fill.fgColor.rgb if c.fill.fill_type else None,
                "border": c.border.left.style,
                "horizontal": c.alignment.horizontal,
            }
            for row in ws
            for c in row
        ],
        "widths": {key: dimension.width for key, dimension in ws.column_dimensions.items()},
    }


@pytest.mark.parametrize("case", GOLDEN["records"], ids=lambda case: case["name"])
def test_observed_flowlib_cells_formats_styles_widths(case):
    rows = copy.deepcopy(case["rows"])
    if case["name"] == "ordinary":
        rows[0][2] = date.fromisoformat(rows[0][2])
    table = ReportTable("Report", case["columns"], rows, start_row=case["start_row"])
    output = render_workbook([table], ExcelOptions(zebra=case["zebra"], border=case["border"]))
    actual = snapshot(load_workbook(io.BytesIO(output)).active)
    expected = copy.deepcopy(case["observed"])
    if case["name"] == "formula_legacy":
        assert expected["cells"][1]["type"] == "f"  # Actual observed unsafe original behavior.
        expected["cells"][1]["type"] = "s"  # Explicitly reviewed technical correction.
    assert actual == expected


@pytest.mark.parametrize("payload", ["=1+1", "+cmd", "-1+2", "@SUM(A1)", "\t=1+1", "#REF!"])
def test_all_strings_are_literal_in_headers_and_cells(payload):
    output = render_workbook([ReportTable("Safe", [payload], [[payload]])])
    ws = load_workbook(io.BytesIO(output), data_only=False).active
    assert ws["A1"].value == ws["A2"].value == payload
    assert ws["A1"].data_type == ws["A2"].data_type == "s"
    with ZipFile(io.BytesIO(output)) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml")
        assert b"<f>" not in xml
        assert not any("externalLink" in name for name in archive.namelist())


def test_multiple_sheets_mappings_date_types_and_explicit_formats():
    rows = [{"ID": "00001", "Betrag": 12.5, "Date": date(2024, 1, 2), "Flag": True}]
    before = copy.deepcopy(rows)
    output = render_workbook(
        [
            ReportTable("Data", ["ID", "Betrag", "Date", "Flag"], iter(rows), formats={"ID": "@"}),
            ReportTable("Plain", ["Betrag"], [[12.5]], profile="plain-v1"),
        ]
    )
    workbook = load_workbook(io.BytesIO(output))
    assert workbook.sheetnames == ["Data", "Plain"]
    assert workbook["Data"]["A2"].value == "00001"
    assert workbook["Data"]["A2"].number_format == "@"
    assert workbook["Data"]["B2"].number_format == get_number_format("Betrag")
    assert workbook["Data"]["C2"].value == datetime(2024, 1, 2)
    assert workbook["Data"]["D2"].value is True
    assert workbook["Data"].freeze_panes == "A2"
    assert workbook["Data"].auto_filter.ref == "A1:D2"
    assert workbook["Plain"]["A2"].number_format == "General"
    assert rows == before


@pytest.mark.parametrize("profile", ["flowlib-legacy-v1", "plain-v1"])
def test_profiles_preserve_existing_api(profile):
    assert get_number_format("Gesamtrate") == '#,##0.00 "EUR"'
    expected = get_number_format("Quote") if profile == "flowlib-legacy-v1" else "General"
    assert get_profile_format(profile, "Quote") == expected
    with pytest.raises(ValueError):
        get_profile_format("jkb", "Quote")


@pytest.mark.parametrize("name", ["", "a" * 32, "A/B", "A\\B", "A?B", "'A", "A'", "a\0b"])
def test_invalid_sheet_names(name):
    with pytest.raises(ValueError):
        render_workbook([ReportTable(name, ["Value"], [])])


def test_duplicate_sheet_names_are_not_silently_renamed():
    with pytest.raises(ValueError, match="Duplicate"):
        render_workbook([ReportTable("Data", ["X"], []), ReportTable("DATA", ["Y"], [])])


@pytest.mark.parametrize(
    "columns,rows",
    [
        ([], []),
        (["A", "A"], []),
        ([""], []),
        (["A"], [[1, 2]]),
        (["A"], [{"B": 1}]),
        (["A"], [{"A": 1, "B": 2}]),
        (["A"], ["text"]),
    ],
)
def test_invalid_shape(columns, rows):
    with pytest.raises((ValueError, TypeError)):
        render_workbook([ReportTable("Data", columns, rows)])


@pytest.mark.parametrize(
    "value",
    [
        float("inf"),
        float("-inf"),
        10**16,
        b"binary",
        [1],
        "\0",
        "\ud800",
        "x" * 32768,
        "😀" * 16384,
        datetime(2024, 1, 1, tzinfo=UTC),
    ],
)
def test_invalid_values_are_rejected_instead_of_truncated_or_corrupted(value):
    with pytest.raises((ValueError, TypeError)):
        render_workbook([ReportTable("Data", ["Value"], [[value]])])


@pytest.mark.parametrize(
    "field,value,tables",
    [
        ("max_rows_per_sheet", 1, [ReportTable("Data", ["A"], [[1], [2]])]),
        ("max_columns", 1, [ReportTable("Data", ["A", "B"], [])]),
        ("max_sheets", 1, [ReportTable("A", ["A"], []), ReportTable("B", ["B"], [])]),
        ("max_cells", 1, [ReportTable("Data", ["A"], [[1]])]),
        ("max_text_characters", 1, [ReportTable("Data", ["A"], [["B"]])]),
        ("max_output_bytes", 1, [ReportTable("Data", ["A"], [])]),
    ],
)
def test_resource_limits(field, value, tables):
    options = ExcelOptions(limits=replace(WorkbookLimits(), **{field: value}))
    with pytest.raises(WorkbookLimitError):
        render_workbook(tables, options)


def test_generator_is_consumed_once_and_stopped_at_limit():
    consumed = []

    def rows():
        for i in range(100):
            consumed.append(i)
            yield [i]

    with pytest.raises(WorkbookLimitError):
        render_workbook(
            [ReportTable("Data", ["Value"], rows())],
            ExcelOptions(limits=WorkbookLimits(max_rows_per_sheet=2)),
        )
    assert consumed == [0, 1, 2]


def test_empty_workbook_invalid_options_and_excel_row_bound():
    for tables, options in [
        ([], ExcelOptions()),
        ([ReportTable("Data", ["A"], [])], ExcelOptions(min_width=100, max_width=10)),
        ([ReportTable("Data", ["A"], [])], ExcelOptions(limits=WorkbookLimits(max_cells=0))),
        ([ReportTable("Data", ["A"], [[1]], start_row=1048576)], ExcelOptions()),
    ]:
        with pytest.raises(ValueError):
            render_workbook(tables, options)


def test_versioned_profile_metadata_and_unknown_profiles():
    from auditcore_reporting import PROFILE_IDS, get_profile_metadata

    for name in PROFILE_IDS:
        metadata = get_profile_metadata(name)
        assert metadata["package_version"] == "0.2.0"
        assert metadata["version"] == "1.0.0"
        assert metadata["status"] == "Draft"
        assert datetime.fromisoformat(metadata["created_at"]).tzinfo is not None
        assert metadata["organization"] == "UNKNOWN"
        assert metadata["content_hash"]
        metadata["title"] = "local mutation"
        assert get_profile_metadata(name)["title"] != "local mutation"
    with pytest.raises(ValueError):
        get_profile_metadata("jkb")


def test_core_profiles_reject_tampered_content(tmp_path, monkeypatch):
    import importlib.resources
    from importlib.resources import files

    from auditcore_reporting import get_profile_metadata

    root = files("auditcore_reporting")
    registry = json.loads(root.joinpath("profile-registry.json").read_text())
    registry["profiles"]["plain-v1"]["content"]["profile_id"] = "modified"
    (tmp_path / "profile-registry.json").write_text(json.dumps(registry))
    monkeypatch.setattr(importlib.resources, "files", lambda name: tmp_path)
    with pytest.raises(ValueError, match="fingerprint"):
        get_profile_metadata("plain-v1")


def test_invalid_option_object_not_silently_ignored():
    with pytest.raises(TypeError):
        render_workbook([], False)


def test_generated_font_properties_follow_office_schema_sequence():
    from defusedxml.ElementTree import fromstring

    payload = render_workbook([ReportTable("Report", ["Betrag"], [[12.5]])])
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    order = (
        "b", "i", "strike", "condense", "extend", "outline", "shadow", "u", "vertAlign",
        "sz", "color", "name", "family", "charset", "scheme",
    )
    with ZipFile(io.BytesIO(payload)) as archive:
        root = fromstring(archive.read("xl/styles.xml"))
        fonts = root.findall(f"{ns}fonts/{ns}font")
        assert len(fonts) == 2
        for font in fonts:
            positions = [order.index(child.tag.removeprefix(ns)) for child in font]
            assert positions == sorted(positions)
    workbook = load_workbook(io.BytesIO(payload))
    assert workbook.active["A1"].font.bold
    assert workbook.active["A1"].font.color.rgb == "00FFFFFF"
    assert workbook.active["A2"].value == 12.5
    workbook.close()
