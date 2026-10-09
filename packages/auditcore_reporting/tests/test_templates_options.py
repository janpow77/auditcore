"""Empty tables and fields, render options (document properties) – issue #235 items 6–8."""

from __future__ import annotations

import io
from datetime import UTC, datetime, timedelta, timezone
from typing import Any

import pytest
from docx import Document
from pypdf import PdfReader

from auditcore_reporting.templates import (
    RenderOptions,
    TemplateError,
    builtin_registry,
    define_template,
    render,
    resolve,
)
from auditcore_reporting.templates.resolve import RFields, RParagraph, RTable

SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {
        "az": {"type": "string"},
        "betreff": {"type": "string"},
        "posten": {
            "type": "array",
            "items": {"type": "object", "properties": {"text": {"type": "string"}}},
        },
    },
}
TABLE: dict[str, Any] = {
    "type": "table",
    "source": "posten",
    "as": "p",
    "empty": "Keine Posten.",
    "columns": [
        {"header": "Nr.", "cell": "1"},
        {"header": "Posten", "cell": "{{ p.text }}"},
    ],
}
FIELDS: dict[str, Any] = {
    "type": "fields",
    "rows": [
        {"label": "Aktenzeichen", "value": "{{ az }}"},
        {"label": "Betreff", "value": "{{ betreff }}"},
    ],
}


def template(*blocks: dict[str, Any]) -> Any:
    return define_template(
        {"id": "optionen", "version": "1.0.0", "title": "Vermerk {{ az }}", "schema": SCHEMA,
         "blocks": list(blocks)}
    )  # fmt: skip


# --- 7. Tabellen ohne Zeilen -------------------------------------------------------------


def test_empty_table_keeps_the_old_default() -> None:
    nodes = resolve(template(TABLE), {"posten": []}).nodes
    assert nodes == (RParagraph("Keine Posten."),)


def test_empty_table_draws_header_on_request() -> None:
    nodes = resolve(template({**TABLE, "header_if_empty": True}), {"posten": []}).nodes
    assert nodes == (RTable(("Nr.", "Posten"), ("left", "left"), ()), RParagraph("Keine Posten."))


@pytest.mark.parametrize("output", ["docx", "pdf", "html"])
def test_empty_table_header_in_every_format(output: str) -> None:
    content = render(template({**TABLE, "header_if_empty": True}), {"posten": []}, output).content
    if output == "docx":
        document = Document(io.BytesIO(content))
        assert [c.text for c in document.tables[0].rows[0].cells] == ["Nr.", "Posten"]
        assert len(document.tables[0].rows) == 1
    elif output == "pdf":
        text = PdfReader(io.BytesIO(content)).pages[0].extract_text()
        assert "Posten" in text and "Keine Posten." in text
    else:
        html = content.decode()
        assert "<thead><tr><th" in html and "<tbody></tbody>" in html


def test_header_if_empty_must_be_boolean() -> None:
    with pytest.raises(TemplateError, match="header_if_empty"):
        template({**TABLE, "header_if_empty": "ja"})


# --- 8. Leere Felder ---------------------------------------------------------------------


def test_empty_field_is_skipped_by_default() -> None:
    nodes = resolve(template(FIELDS), {"az": "A-1", "betreff": " "}).nodes
    assert nodes == (RFields((("Aktenzeichen", "A-1"),)),)
    assert resolve(template(FIELDS), {}).nodes == ()


def test_empty_field_shows_the_configured_placeholder() -> None:
    nodes = resolve(template({**FIELDS, "empty": "—"}), {"az": "A-1"}).nodes
    assert nodes == (RFields((("Aktenzeichen", "A-1"), ("Betreff", "—"))),)
    assert resolve(template({**FIELDS, "empty": ""}), {}).nodes == (
        RFields((("Aktenzeichen", ""), ("Betreff", ""))),
    )


@pytest.mark.parametrize("output", ["docx", "pdf", "html"])
def test_empty_field_placeholder_in_every_format(output: str) -> None:
    content = render(template({**FIELDS, "empty": "—"}), {"az": "A-1"}, output).content
    if output == "docx":
        cells = [c.text for c in Document(io.BytesIO(content)).tables[0].rows[1].cells]
        assert cells == ["Betreff", "—"]
    elif output == "pdf":
        assert "Betreff\n—" in PdfReader(io.BytesIO(content)).pages[0].extract_text()
    else:
        assert "<td>Betreff</td><td>—</td>" in content.decode()


def test_empty_placeholder_is_checked_against_the_contract() -> None:
    with pytest.raises(TemplateError, match="unbekannt"):
        template({**FIELDS, "empty": "{{ unbekannt }}"})
    with pytest.raises(TemplateError, match="empty"):
        template({**FIELDS, "empty": 0})


# --- 6. Dokumenteigenschaften ------------------------------------------------------------

CREATED = datetime(2026, 10, 4, 12, 30, 15, tzinfo=timezone(timedelta(hours=2)))
OPTIONS = RenderOptions(author="Prüferin Muster", title="Vermerk Musterakte", created=CREATED)


def test_docx_core_properties_from_options() -> None:
    content = render(template(FIELDS), {"az": "A-1"}, "docx", options=OPTIONS).content
    core = Document(io.BytesIO(content)).core_properties
    assert core.author == "Prüferin Muster" and core.last_modified_by == "Prüferin Muster"
    assert core.title == "Vermerk Musterakte"
    assert core.created == datetime(2026, 10, 4, 10, 30, 15)
    assert core.modified == core.created
    assert core.language == "de-DE" and "Vorlage optionen 1.0.0" in core.comments


def test_docx_without_options_has_no_author_or_time() -> None:
    content = render(template(FIELDS), {"az": "A-1"}, "docx").content
    core = Document(io.BytesIO(content)).core_properties
    assert core.author == "" and core.created is None and core.title == "Vermerk A-1"


def test_pdf_info_from_options() -> None:
    content = render(template(FIELDS), {"az": "A-1"}, "pdf", options=OPTIONS).content
    info = PdfReader(io.BytesIO(content)).metadata
    assert info is not None
    assert info.author == "Prüferin Muster" and info.title == "Vermerk Musterakte"
    assert info.creation_date == datetime(2026, 10, 4, 10, 30, 15, tzinfo=UTC)
    assert info.modification_date == info.creation_date


def test_html_meta_from_options() -> None:
    html = render(template(FIELDS), {"az": "A-1"}, "html", options=OPTIONS).content.decode()
    assert '<meta name="author" content="Prüferin Muster">' in html
    assert '<meta name="dcterms.created" content="2026-10-04T10:30:15Z">' in html
    assert "<title>Vermerk Musterakte</title>" in html


@pytest.mark.parametrize("output", ["docx", "pdf", "html"])
def test_options_are_deterministic(output: str) -> None:
    first = render(template(FIELDS), {"az": "A-1"}, output, options=OPTIONS).content
    second = render(template(FIELDS), {"az": "A-1"}, output, options=OPTIONS).content
    assert first == second


def test_options_are_validated() -> None:
    with pytest.raises(TemplateError, match="Zeitzone"):
        RenderOptions(created=datetime(2026, 10, 4))
    with pytest.raises(TemplateError, match="author"):
        RenderOptions(author="a\nb")
    with pytest.raises(TemplateError, match="title"):
        RenderOptions(title="x" * 256)


def test_docx_template_rejects_options() -> None:
    document = Document()
    document.add_paragraph("Az. {{ az }}")
    buffer = io.BytesIO()
    document.save(buffer)
    word = define_template(
        {"id": "wort", "version": "1.0.0", "title": "Wort", "schema": SCHEMA},
        docx=buffer.getvalue(),
    )
    assert render(word, {"az": "A-1"}, "docx").content
    with pytest.raises(TemplateError, match="Render-Optionen"):
        render(word, {"az": "A-1"}, "docx", options=OPTIONS)


@pytest.mark.parametrize("output", ["docx", "pdf", "html"])
def test_default_options_keep_builtin_output(output: str) -> None:
    for memo in (builtin_registry().get("vermerk"), builtin_registry().get("pruefbericht")):
        plain = render(memo, memo.sample, output).content
        assert render(memo, memo.sample, output, options=RenderOptions()).content == plain
