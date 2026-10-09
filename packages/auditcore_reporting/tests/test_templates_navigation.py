"""Pictures, jump targets, links, table of contents and PDF outline – issue #235 items 1 and 5."""

from __future__ import annotations

import base64
import io
import zipfile
from typing import Any

import pytest
from docx import Document
from PIL import Image as PilImage
from pypdf import PdfReader

from auditcore_reporting.templates import (
    DesignProfile,
    RenderLimitError,
    RenderOptions,
    ReportImage,
    ResolveLimits,
    TemplateError,
    builtin_registry,
    define_template,
    design_from_dict,
    render,
    resolve,
)
from auditcore_reporting.templates.resolve import RContents, RHeading, RImage, RParagraph

SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {
        "foto": {"type": "string"},
        "zahl": {"type": "number"},
        "kapitel": {
            "type": "array",
            "items": {"type": "object", "properties": {"titel": {"type": "string"}}},
        },
    },
}
CHAPTERS = {"kapitel": [{"titel": "Erstes"}, {"titel": "Zweites"}, {"titel": "Drittes"}]}


def picture(kind: str = "PNG", size: tuple[int, int] = (40, 20)) -> bytes:
    buffer = io.BytesIO()
    PilImage.new("RGB", size, (200, 30, 30)).save(buffer, kind)
    return buffer.getvalue()


def template(*blocks: dict[str, Any], **extra: object) -> Any:
    return define_template(
        {"id": "akte", "version": "1.0.0", "title": "Akte", "schema": SCHEMA,
         "blocks": list(blocks), **extra}
    )  # fmt: skip


def docx_xml(content: bytes, part: str = "word/document.xml") -> str:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return archive.read(part).decode()


# --- 5. Bilder -----------------------------------------------------------------------------

WAPPEN = ReportImage("wappen", picture())
LOGO = ReportImage("logo", picture("JPEG", (30, 30)))
IMAGES = template(
    {"type": "image", "image": "wappen", "width_cm": 4, "align": "center", "alt": "Wappen"},
    {"type": "image", "image": "logo"},
    {"type": "image", "source": "foto"},
)
OPTIONS = RenderOptions(images=(WAPPEN,))
DESIGN = DesignProfile(id="amt", images=(LOGO, ReportImage("wappen", picture(size=(10, 10)))))


def test_report_image_reads_format_and_size() -> None:
    assert (WAPPEN.kind, WAPPEN.width_px, WAPPEN.height_px) == ("png", 40, 20)
    assert (LOGO.kind, LOGO.width_px, LOGO.height_px) == ("jpeg", 30, 30)
    for data in (b"GIF89a....", b"\x89PNG\r\n\x1a\n", b"\xff\xd8\xff\xe0"):
        with pytest.raises(TemplateError, match="PNG oder JPEG"):
            ReportImage("x", data)
    with pytest.raises(TemplateError, match="Name"):
        ReportImage("../wappen", picture())


def test_images_resolve_from_options_design_and_data() -> None:
    foto = "data:image/png;base64," + base64.b64encode(picture(size=(8, 4))).decode()
    nodes = resolve(IMAGES, {"foto": foto}, images=(*OPTIONS.images, *DESIGN.images)).nodes
    assert [type(n) for n in nodes] == [RImage, RImage, RImage]
    first, second, third = nodes
    assert isinstance(first, RImage) and first.image is WAPPEN  # options before design
    assert first.size_cm(16.0) == (4.0, 2.0) and first.alt == "Wappen"
    assert isinstance(second, RImage) and second.image is LOGO
    assert isinstance(third, RImage) and third.image.width_px == 8
    assert third.size_cm(0.1) == (0.1, 0.05)  # never wider than the text


def test_missing_or_empty_pictures() -> None:
    with pytest.raises(TemplateError, match="fehlt"):
        resolve(IMAGES, {}, images=(LOGO,))
    nodes = resolve(IMAGES, {"foto": ""}, images=(WAPPEN, LOGO)).nodes
    assert len(nodes) == 2  # an empty data picture is skipped like an empty field
    with pytest.raises(TemplateError, match="Base64"):
        resolve(IMAGES, {"foto": "kein bild!"}, images=(WAPPEN, LOGO))
    with pytest.raises(RenderLimitError, match="Bilder"):
        resolve(IMAGES, {}, ResolveLimits(max_image_bytes=10), images=(WAPPEN, LOGO))


def test_image_definition_is_checked() -> None:
    with pytest.raises(TemplateError, match="genau eins"):
        template({"type": "image"})
    with pytest.raises(TemplateError, match="Text"):
        template({"type": "image", "source": "zahl"})
    with pytest.raises(TemplateError, match="width_cm"):
        template({"type": "image", "image": "x", "width_cm": 31})
    with pytest.raises(TemplateError, match="nicht deklariert"):
        template({"type": "image", "source": "pfad.zur.datei"})


def test_images_in_docx() -> None:
    content = render(IMAGES, {}, "docx", DESIGN, options=OPTIONS).content
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        names = archive.namelist()
        assert archive.read("word/media/image1.png") == WAPPEN.data
        assert archive.read("word/media/image2.jpeg") == LOGO.data
    assert "word/media/image3.png" not in names
    types = docx_xml(content, "[Content_Types].xml")
    assert 'Extension="jpeg" ContentType="image/jpeg"' in types
    assert 'Extension="png" ContentType="image/png"' in types
    document = Document(io.BytesIO(content))
    shapes = document.inline_shapes
    assert len(shapes) == 2
    assert shapes[0].width == 4 * 360_000 and shapes[0].height == 2 * 360_000
    assert 'descr="Wappen"' in docx_xml(content)


def test_images_in_pdf_and_html() -> None:
    pdf = render(IMAGES, {}, "pdf", DESIGN, options=OPTIONS).content
    assert len(PdfReader(io.BytesIO(pdf)).pages[0].images) == 2
    html = render(IMAGES, {}, "html", DESIGN, options=OPTIONS).content.decode()
    assert "img-src data:" in html and html.count("<img ") == 2
    assert 'src="data:image/png;base64,' in html and 'alt="Wappen"' in html
    plain = render(
        builtin_registry().get("vermerk"), builtin_registry().get("vermerk").sample, "html"
    )
    assert "img-src" not in plain.content.decode()


def test_images_are_python_only_in_design_profiles() -> None:
    assert DESIGN.to_dict()["images"] == ["logo", "wappen"]
    with pytest.raises(TemplateError, match="Python-API"):
        design_from_dict({"images": ["logo"]})


# --- 1. Sprungmarken, Links, Inhaltsverzeichnis, Lesezeichen -------------------------------

NAVIGATION = template(
    {"type": "toc", "title": "Inhalt", "levels": 2},
    {"type": "paragraph", "text": "Zum Anhang", "link": "anhang"},
    {"type": "section", "title": "{{ k.titel }}", "for": "kapitel", "as": "k", "level": 1,
     "anchor": "kapitel", "blocks": [
         {"type": "heading", "text": "Unterpunkt", "level": 2},
         {"type": "heading", "text": "Detail", "level": 3},
         {"type": "pagebreak"}]},
    {"type": "heading", "text": "Anhang", "anchor": "anhang"},
    outline=True,
)  # fmt: skip


def test_targets_and_contents_resolve() -> None:
    nodes = resolve(NAVIGATION, CHAPTERS).nodes
    headings = [(n.text, n.anchor) for n in nodes if isinstance(n, RHeading)]
    assert headings[:4] == [
        ("Erstes", "kapitel"), ("Unterpunkt", "auto-1"), ("Detail", "auto-2"),
        ("Zweites", "kapitel-2"),
    ]  # fmt: skip
    assert headings[-1] == ("Anhang", "anhang")
    contents = nodes[0]
    assert isinstance(contents, RContents)
    assert all(level <= 2 for level, _, _ in contents.entries) and len(contents.entries) == 7
    assert nodes[1] == RParagraph("Zum Anhang", "", "anhang")


def test_links_must_point_to_declared_anchors() -> None:
    with pytest.raises(TemplateError, match="nicht deklariert"):
        template({"type": "paragraph", "text": "x", "link": "fehlt"})
    with pytest.raises(TemplateError, match="mehrfach"):
        template(
            {"type": "heading", "text": "a", "anchor": "x"},
            {"type": "heading", "text": "b", "anchor": "x"},
        )
    with pytest.raises(TemplateError, match="Sprungmarke"):
        template({"type": "heading", "text": "a", "anchor": "auto-1"})
    with pytest.raises(TemplateError, match="levels"):
        template({"type": "toc", "levels": 4})


def test_pdf_outline_contents_and_links_jump_to_the_right_page() -> None:
    reader = PdfReader(io.BytesIO(render(NAVIGATION, CHAPTERS, "pdf").content))
    outline = reader.outline
    top = [e for e in outline if isinstance(e, dict)]
    assert [entry["/Title"] for entry in top] == ["Erstes", "Zweites", "Drittes", "Anhang"]
    nested = outline[1]
    assert isinstance(nested, list) and nested[0]["/Title"] == "Unterpunkt"
    assert isinstance(nested[1], list) and nested[1][0]["/Title"] == "Detail"
    pages = [reader.get_destination_page_number(entry) for entry in top]
    assert pages == [0, 1, 2, 3]  # Erstes, Zweites, Drittes, Anhang (contents on page 1)
    first_page = reader.pages[0].extract_text()
    assert "Inhalt" in first_page and "Zweites" in first_page
    assert first_page.count("Detail") == 1  # body only: level 3 is not in the contents
    assert "Seite 1 von 4" in first_page
    targets = [
        annotation.get_object()["/Dest"][0] for annotation in reader.pages[0].get("/Annots", [])
    ]
    page_ids = [page.indirect_reference for page in reader.pages]
    assert page_ids[3] in targets  # the link "Zum Anhang"
    assert (
        render(NAVIGATION, CHAPTERS, "pdf").content == render(NAVIGATION, CHAPTERS, "pdf").content
    )


def test_docx_bookmarks_links_and_contents() -> None:
    xml = docx_xml(render(NAVIGATION, CHAPTERS, "docx").content)
    assert '<w:bookmarkStart w:id="1" w:name="ac_kapitel"/>' in xml
    assert 'w:name="ac_kapitel_2"' in xml and 'w:name="ac_anhang"' in xml
    assert '<w:hyperlink w:anchor="ac_anhang" w:history="1">' in xml
    assert xml.count('<w:hyperlink w:anchor="ac_auto_1"') == 1  # contents line
    assert 'w:anchor="ac_auto_2"' not in xml  # level 3 is not in the contents
    paragraphs = [
        p.text
        for p in Document(io.BytesIO(render(NAVIGATION, CHAPTERS, "docx").content)).paragraphs
    ]
    assert paragraphs[:3] == ["Inhalt", "Erstes", "Unterpunkt"]


def test_html_ids_links_and_contents() -> None:
    html = render(NAVIGATION, CHAPTERS, "html").content.decode()
    assert '<h1 id="kapitel">Erstes</h1>' in html and '<h1 id="kapitel-2">Zweites</h1>' in html
    assert '<a href="#anhang">Zum Anhang</a>' in html
    assert "<nav><p><b>Inhalt</b></p>" in html and 'href="#auto-1"' in html


def test_documents_without_navigation_keep_their_output() -> None:
    plain = template({"type": "heading", "text": "Titel"}, {"type": "paragraph", "text": "Text"})
    nodes = resolve(plain, {}).nodes
    assert nodes == (RHeading("Titel", 1), RParagraph("Text"))
    xml = docx_xml(render(plain, {}, "docx").content)
    assert "bookmark" not in xml and "hyperlink" not in xml
    outlined = template({"type": "heading", "text": "Titel"}, outline=True)
    assert resolve(outlined, {}).nodes == (RHeading("Titel", 1, "auto-1"),)
    reader = PdfReader(io.BytesIO(render(outlined, {}, "pdf").content))
    assert [entry["/Title"] for entry in reader.outline] == ["Titel"]  # type: ignore[index]
