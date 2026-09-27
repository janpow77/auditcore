"""Rendering of structured templates: identical content in DOCX, PDF and HTML, byte-stable."""

from __future__ import annotations

import hashlib
import io
import subprocess
import sys
import zipfile
from typing import Any

import pytest
from docx import Document
from pypdf import PdfReader

from auditcore_reporting.templates import (
    NEUTRAL_DESIGN,
    RenderLimitError,
    ResolveLimits,
    TemplateError,
    builtin_registry,
    design_from_dict,
    render,
)

REGISTRY = builtin_registry()
REPORT = REGISTRY.get("pruefbericht")
MEMO = REGISTRY.get("vermerk")

#: Golden digests: HTML bytes and the uncompressed DOCX parts (name + SHA-256 of
#: each part, so the digest does not depend on the zlib build). A change here is a
#: visible change of the output and needs a new template or design version.
GOLDEN = {
    ("pruefbericht", "html"): "814b0b8ba66a81e1b9eb11fbb9c5808b97092ef2f85304d4124e734da07e63db",
    ("pruefbericht", "docx"): "d6043c56e6352e35a006a9aa3649e5c4ccad98841926b64fc1bb56bb356c7d0a",
    ("vermerk", "html"): "950fa91bdc74d21836e9c2da2e7ba0619dbdcbabc8c2d1a72cd581e9d67eaef3",
    ("vermerk", "docx"): "304c715de6842102142de3df7f96b4b7ed19533b4b7513f42f2d99e5194ff565",
}


def parts_digest(content: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        lines = [
            f"{info.filename}:{hashlib.sha256(archive.read(info)).hexdigest()}"
            for info in archive.infolist()
        ]
    return hashlib.sha256("\n".join(lines).encode()).hexdigest()


@pytest.mark.parametrize("template_id", ["pruefbericht", "vermerk"])
@pytest.mark.parametrize("output", ["docx", "html"])
def test_output_matches_golden_digest(template_id: str, output: str) -> None:
    template = REGISTRY.get(template_id)
    content = render(template, template.sample, output).content
    digest = parts_digest(content) if output == "docx" else hashlib.sha256(content).hexdigest()
    assert digest == GOLDEN[(template_id, output)]


@pytest.mark.parametrize("output", ["docx", "pdf", "html"])
def test_rendering_is_byte_identical(output: str) -> None:
    first = render(REPORT, REPORT.sample, output)
    second = render(REPORT, dict(REPORT.sample), output)
    assert first.content == second.content
    assert first.template_fingerprint == REPORT.fingerprint
    assert first.data_sha256 == second.data_sha256 and len(first.data_sha256) == 64
    assert first.design == "neutral-v1@1.0.0"


def test_docx_is_a_clean_word_document() -> None:
    content = render(REPORT, REPORT.sample, "docx").content
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        names = archive.namelist()
        assert all(info.date_time == (1980, 1, 1, 0, 0, 0) for info in archive.infolist())
        core = archive.read("docProps/core.xml").decode()
    assert not any("vba" in name.lower() or name.endswith(".bin") for name in names)
    assert "Vorlage pruefbericht 1.0.0" in core
    document = Document(io.BytesIO(content))
    styles = [(p.style.name, p.text) for p in document.paragraphs if p.text]
    assert styles[0] == ("Heading 1", "Prüfbericht")
    assert ("Heading 3", "Feststellung 2: Rechnung außerhalb des Förderzeitraums") in styles
    texts = [text for _, text in styles]
    assert (
        "Die Ergebnisse der Verwaltungsüberprüfung nach Art. 74 der Verordnung (EU) 2021/1060 wurden in die Prüfung einbezogen."
        in texts
    )
    table = document.tables[1]
    assert [c.text for c in table.rows[2].cells] == [
        "2",
        "Rechnung außerhalb des Förderzeitraums",
        "finanziell",
        "1.234,56 €",
    ]


def test_pdf_carries_the_same_text() -> None:
    content = render(REPORT, REPORT.sample, "pdf").content
    reader = PdfReader(io.BytesIO(content))
    text = "\n".join(page.extract_text() for page in reader.pages)
    for expected in ("Prüfbericht", "Feststellung 2", "1.234,56 €", "Verwaltungsüberprüfung"):
        assert expected in text
    assert "Seite 1 von" in text
    assert reader.metadata is not None and reader.metadata.title == "Prüfbericht PB-2026-0042"


def test_text_blocks_follow_the_audit_result() -> None:
    data: dict[str, Any] = {**REPORT.sample, "feststellungen": [], "summe_nicht_foerderfaehig": 0}
    data["verwaltungsueberpruefung_einbezogen"] = False
    result = render(REPORT, data, "html")
    html = result.content.decode()
    assert result.text_blocks == ("rechtsgrundlage", "ohne_feststellungen")
    assert "Die Prüfung hat zu keinen Feststellungen geführt." in html
    assert "Stellungnahme" not in html and "<table><thead>" not in html


def test_data_are_escaped_everywhere() -> None:
    data = {**MEMO.sample, "anlass": '<script>alert("x")</script> & „Zitat“'}
    html = render(MEMO, data, "html").content.decode()
    assert "<script>" not in html and "&lt;script&gt;" in html
    docx = render(MEMO, data, "docx").content
    assert any(p.text.startswith("<script>") for p in Document(io.BytesIO(docx)).paragraphs)
    pdf = render(MEMO, data, "pdf").content
    assert "<script>" in PdfReader(io.BytesIO(pdf)).pages[0].extract_text()


def test_design_profile_is_exchangeable() -> None:
    design = design_from_dict(
        {
            "id": "amt-v1",
            "accent_color": "7A1F1F",
            "header_text": "Musteramt",
            "footer_text": "Intern",
        }
    )
    neutral = render(MEMO, MEMO.sample, "docx", NEUTRAL_DESIGN).content
    custom = render(MEMO, MEMO.sample, "docx", design).content
    assert neutral != custom
    with zipfile.ZipFile(io.BytesIO(custom)) as archive:
        assert "Musteramt" in archive.read("word/header1.xml").decode()
        assert "7A1F1F" in archive.read("word/styles.xml").decode()
    assert "Musteramt" in render(MEMO, MEMO.sample, "html", design).content.decode()


def test_limits_and_formats() -> None:
    with pytest.raises(RenderLimitError):
        render(REPORT, REPORT.sample, "docx", limits=ResolveLimits(max_loop_items=1))
    with pytest.raises(RenderLimitError):
        render(REPORT, REPORT.sample, "html", limits=ResolveLimits(max_characters=100))
    with pytest.raises(TemplateError, match="Format"):
        render(REPORT, REPORT.sample, "odt")


def test_pdf_without_extra_reports_dependency() -> None:
    script = """
import importlib.abc, sys
class Block(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, *args):
        if fullname.split(".")[0] == "reportlab":
            raise ModuleNotFoundError("blocked", name=fullname)
sys.meta_path.insert(0, Block())
from auditcore_reporting.templates import RenderDependencyError, builtin_registry, render
template = builtin_registry().get("vermerk")
assert render(template, template.sample, "docx").content[:2] == b"PK"
try:
    render(template, template.sample, "pdf")
except RenderDependencyError as exc:
    assert "auditcore_reporting[pdf]" in str(exc)
else:
    raise AssertionError("expected RenderDependencyError")
"""
    result = subprocess.run([sys.executable, "-I", "-c", script], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
