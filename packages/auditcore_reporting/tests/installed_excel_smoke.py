"""Execute with python -I after installing the excel extra; real XLSX readback."""

from __future__ import annotations

import io
from datetime import date
from importlib.metadata import version
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

import auditcore_reporting
from auditcore_reporting import ExcelOptions, ReportTable, render_workbook


def main() -> None:
    """Verify actual installed bytes, formats, typed values and inert formula-like text."""
    assert version("auditcore_reporting") == "0.3.0"
    location = Path(auditcore_reporting.__file__).resolve()
    assert "site-packages" in location.parts or "dist-packages" in location.parts, location
    output = render_workbook(
        [
            ReportTable(
                "Prüfung",
                ["Betrag", "Quote", "Datum", "Text"],
                [[123.45, 0.25, date(2024, 1, 2), "=1+1"], [None, 0.5, date(2024, 2, 3), "ÄÖÜ"]],
            ),
            ReportTable("Neutral", ["Betrag"], [[10]], profile="plain-v1"),
        ],
        ExcelOptions(),
    )
    workbook = load_workbook(io.BytesIO(output), data_only=False)
    assert workbook.sheetnames == ["Prüfung", "Neutral"]
    ws = workbook["Prüfung"]
    assert ws["A2"].value == 123.45 and ws["A2"].number_format == '#,##0.00 "EUR"'
    assert ws["B2"].number_format == "0.00%"
    assert ws["C2"].value.date() == date(2024, 1, 2)
    assert ws["D2"].value == "=1+1" and ws["D2"].data_type == "s"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:D3"
    assert workbook["Neutral"]["A2"].number_format == "General"
    with ZipFile(io.BytesIO(output)) as archive:
        assert b"<f>" not in archive.read("xl/worksheets/sheet1.xml")
    print(
        "PASS: installed Excel renderer, values/formats/injection checks; "
        f"openpyxl {version('openpyxl')}"
    )


if __name__ == "__main__":
    main()
