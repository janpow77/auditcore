"""Word templates: placeholders across runs, control tags, security check of the package."""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable
from typing import Any

import pytest
from docx import Document

from auditcore_reporting.templates import (
    TemplateError,
    UnsafeDocumentError,
    define_template,
    render,
)
from auditcore_reporting.templates._docx_package import MAIN_DOCUMENT, MAIN_TEMPLATE
from auditcore_reporting.templates.docx_template import docx_paragraphs

SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {
        "aktenzeichen": {"type": "string"},
        "empfaenger": {"type": "string"},
        "hinweis": {"type": "string"},
        "eilig": {"type": "boolean"},
        "frist": {"type": "string", "format": "date"},
        "positionen": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {"name": {"type": "string"}, "betrag": {"type": "number"}},
            },
        },
    },
}
TEXT_BLOCKS = [
    {
        "id": "frist",
        "if": "frist",
        "text": "Antwort bis {{ frist | datum }}.\nMit freundlichen Grüßen",
    }
]
DATA = {
    "aktenzeichen": "AZ-1",
    "empfaenger": "Beispiel GmbH",
    "eilig": True,
    "frist": "2026-10-01",
    "positionen": [{"name": "Prüfung Ä", "betrag": 1.5}, {"name": "B & <C>", "betrag": 2}],
}


def word_template(extra: Callable[[Any], None] | None = None) -> bytes:
    """A real python-docx document; Word splits placeholders over runs like this."""
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Az. {{ aktenzeichen }}"
    paragraph = document.add_paragraph()
    paragraph.add_run("Sehr geehrte ")
    paragraph.add_run("{{ empf")
    paragraph.add_run("aenger }}").bold = True
    paragraph.add_run(",")
    for text in (
        "{%p if eilig %}",
        "Eilt!",
        "{%p endif %}",
        "{%p if not hinweis %}",
        "Kein Hinweis.",
    ):
        document.add_paragraph(text)
    document.add_paragraph("{%p endif %}")
    document.add_paragraph("{{ textbaustein.frist }}")
    table = document.add_table(rows=4, cols=2)
    for row, cells in enumerate(
        [("Position", "Betrag"), ("{%tr for p in positionen %}", ""),
         ("{{ p.name }}", "{{ p.betrag | eur }}"), ("{%tr endfor %}", "")]
    ):  # fmt: skip
        for column, text in enumerate(cells):
            table.rows[row].cells[column].text = text
    for text in ("{%p for p in positionen %}", "– {{ p.name }}", "{%p endfor %}"):
        document.add_paragraph(text)
    if extra is not None:
        extra(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def template(docx: bytes, **change: object) -> Any:
    definition: dict[str, Any] = {
        "id": "schreiben",
        "version": "1.0.0",
        "title": "Schreiben",
        "schema": SCHEMA,
        "text_blocks": TEXT_BLOCKS,
        "sample": {},
    }
    definition.update(change)
    return define_template(definition, docx=docx)


def test_placeholders_tags_and_row_loops_are_filled() -> None:
    filled = render(template(word_template()), DATA, "docx")
    document = Document(io.BytesIO(filled.content))
    texts = [p.text for p in document.paragraphs]
    assert texts == [
        "Sehr geehrte Beispiel GmbH,",
        "Eilt!",
        "Kein Hinweis.",
        "Antwort bis 01.10.2026.\nMit freundlichen Grüßen",
        "– Prüfung Ä",
        "– B & <C>",
    ]
    runs = [(run.text, run.bold) for run in document.paragraphs[0].runs]
    assert runs == [("Sehr geehrte ", None), ("Beispiel GmbH", None), ("", True), (",", None)]
    assert [[c.text for c in row.cells] for row in document.tables[0].rows] == [
        ["Position", "Betrag"],
        ["Prüfung Ä", "1,50 €"],
        ["B & <C>", "2,00 €"],
    ]
    assert document.sections[0].header.paragraphs[0].text == "Az. AZ-1"
    assert filled.content == render(template(word_template()), DATA, "docx").content


def test_empty_loops_and_false_conditions_drop_content() -> None:
    data = {"empfaenger": "X", "eilig": False, "hinweis": "da", "positionen": []}
    content = render(template(word_template()), data, "docx").content
    assert [p for p in docx_paragraphs(content) if p] == ["Sehr geehrte X,", "Position", "Betrag"]
    rows = Document(io.BytesIO(content)).tables[0].rows
    assert len(rows) == 1


def test_dotx_becomes_docx() -> None:
    source = rewrite(
        word_template(),
        "[Content_Types].xml",
        lambda xml: xml.replace(MAIN_DOCUMENT.encode(), MAIN_TEMPLATE.encode()),
    )
    content = render(template(source), DATA, "docx").content
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        types = archive.read("[Content_Types].xml").decode()
    assert MAIN_DOCUMENT in types and MAIN_TEMPLATE not in types


@pytest.mark.parametrize(
    ("paragraphs", "message"),
    [
        (["{{ unbekannt }}"], "nicht deklariert"),
        (["{%p if eilig %}"], "ohne Ende-Tag"),
        (["{%p endif %}"], "ohne öffnenden Tag"),
        (["{%p for p in aktenzeichen %}", "{%p endfor %}"], "keine Liste"),
        (["{%p if eilig %}", "{%p endfor %}"], "passt nicht"),
        (["Text {%p if eilig %}"], "allein im Absatz"),
        (["{%p while x %}"], "unbekannter Steuer-Tag"),
        (["{%tr if eilig %}"], "Tabellenzeile"),
        (["{{ p.name }}"], "nicht deklariert"),
    ],
)
def test_template_errors_are_found_at_definition(paragraphs: list[str], message: str) -> None:
    def add(document: Any) -> None:
        for text in paragraphs:
            document.add_paragraph(text)

    with pytest.raises(TemplateError, match=message):
        template(word_template(add))


def test_docx_template_renders_docx_only() -> None:
    with pytest.raises(TemplateError, match="formats"):
        template(word_template(), formats=["pdf"])
    with pytest.raises(TemplateError, match="schließen sich aus"):
        template(word_template(), blocks=[])


def rewrite(source: bytes, name: str, change: Callable[[bytes], bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as src, zipfile.ZipFile(buffer, "w") as dst:
        for info in src.infolist():
            data = src.read(info)
            dst.writestr(info, change(data) if info.filename == name else data)
    return buffer.getvalue()


def add_entry(source: bytes, name: str, data: bytes, compress: bool = True) -> bytes:
    buffer = io.BytesIO(source)
    with zipfile.ZipFile(
        buffer, "a", zipfile.ZIP_DEFLATED if compress else zipfile.ZIP_STORED
    ) as archive:
        archive.writestr(name, data)
    return buffer.getvalue()


def _document_xml(transform: Callable[[str], str]) -> Callable[[bytes], bytes]:
    return lambda raw: transform(raw.decode("utf-8")).encode("utf-8")


RELS = "word/_rels/document.xml.rels"
ATTACHED = (
    '<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
    'relationships/attachedTemplate" Target="file:///C:/vorlage.dotm" TargetMode="External"/>'
)
LINKED_IMAGE = (
    '<Relationship Id="rId98" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
    'relationships/image" Target="http://example.invalid/x.png" TargetMode="External"/>'
)
FIELD = '<w:p><w:r><w:instrText xml:space="preserve"> INCLUDETEXT "c:/x.docx" </w:instrText></w:r></w:p>'
SPLIT_FIELD = "<w:p><w:r><w:instrText> DD</w:instrText></w:r><w:r><w:instrText>EAUTO x </w:instrText></w:r></w:p>"
XXE = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>'


@pytest.mark.parametrize(
    ("mutate", "message"),
    [
        (lambda d: add_entry(d, "word/vbaProject.bin", b"x"), "Makros"),
        (lambda d: add_entry(d, "word/activeX/activeX1.xml", b"<x/>"), "ActiveX"),
        (lambda d: add_entry(d, "../evil.txt", b"x"), "Unsicherer Eintragsname"),
        (lambda d: add_entry(d, "customXml/item9.xml", XXE), "unsicheres XML"),
        (
            lambda d: rewrite(d, "[Content_Types].xml", lambda x: x.replace(
                MAIN_DOCUMENT.encode(),
                b"application/vnd.ms-word.document.macroEnabled.main+xml")),
            "Inhaltstyp",
        ),
        (lambda d: rewrite(d, RELS, lambda x: x.replace(b"</Relationships>", ATTACHED.encode() + b"</Relationships>")), "attachedtemplate"),
        (lambda d: rewrite(d, RELS, lambda x: x.replace(b"</Relationships>", LINKED_IMAGE.encode() + b"</Relationships>")), "externe Quelle"),
        (lambda d: rewrite(d, "word/document.xml", _document_xml(lambda x: x.replace("<w:sectPr", FIELD + "<w:sectPr", 1))), "INCLUDETEXT"),
        (lambda d: rewrite(d, "word/document.xml", _document_xml(lambda x: x.replace("<w:sectPr", SPLIT_FIELD + "<w:sectPr", 1))), "DDEAUTO"),
        (lambda d: rewrite(d, "word/document.xml", _document_xml(lambda x: x.replace("<w:sectPr", '<w:altChunk r:id="rId1"/><w:sectPr', 1))), "eingebettete"),
        (lambda d: add_entry(d, "word/media/bomb.xml", b"<a>" + b" " * (8 * 1024 * 1024) + b"</a>"), "Kompressionsrate"),
        (lambda d: b"kein zip", "keine lesbare|Keine lesbare"),
    ],
)  # fmt: skip
def test_unsafe_packages_are_rejected(mutate: Callable[[bytes], bytes], message: str) -> None:
    with pytest.raises(UnsafeDocumentError, match=message):
        template(mutate(word_template()))


def test_injected_placeholders_in_data_are_not_expanded() -> None:
    data = {**DATA, "empfaenger": "{{ aktenzeichen }} {%p if eilig %}"}
    content = render(template(word_template()), data, "docx").content
    assert docx_paragraphs(content)[0] == "Sehr geehrte {{ aktenzeichen }} {%p if eilig %},"
