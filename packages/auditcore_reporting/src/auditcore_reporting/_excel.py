"""Optional XLSX adapter adapted from MIT Flowlib report/styles/auto-fit rules.

Source: janpow77/flowlib@aca2dc6aad25aea0720312dbcc6da00b0bcba330.
Copyright (c) 2026 Jan Riener. MIT; see LICENSE and NOTICE.
"""

from __future__ import annotations

import math
import re
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import fields
from datetime import date, datetime
from io import BytesIO
from typing import cast

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from auditcore_reporting._ooxml import normalize_font_sequence
from auditcore_reporting.profiles import PROFILE_IDS, get_profile_format
from auditcore_reporting.workbook import (
    CellValue,
    ExcelOptions,
    ReportRow,
    ReportTable,
    WorkbookLimitError,
)

_XML_INVALID = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_SHEET_INVALID = re.compile(r"[\\/*?:\[\]]")
_HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
_ZEBRA = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
_SIDE = Side(style="thin", color="000000")
_BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)


class _Budget:
    def __init__(self, options: ExcelOptions) -> None:
        self.options = options
        self.cells = 0
        self.text = 0

    def consume(self, value: CellValue) -> None:
        """Count each written cell and string before accepting it into the workbook."""
        self.cells += 1
        if isinstance(value, str):
            self.text += len(value)
        if self.cells > self.options.limits.max_cells:
            raise WorkbookLimitError("Workbook cell budget exceeded")
        if self.text > self.options.limits.max_text_characters:
            raise WorkbookLimitError("Workbook text budget exceeded")


def _validate_options(options: ExcelOptions) -> None:
    for descriptor in fields(options.limits):
        value = getattr(options.limits, descriptor.name)
        if isinstance(value, bool) or not isinstance(value, int) or value < 1:
            raise ValueError(f"{descriptor.name} must be a positive integer")
    for flag in (options.zebra, options.border, options.freeze_header, options.auto_filter):
        if not isinstance(flag, bool):
            raise TypeError("Formatting flags must be booleans")
    if not (math.isfinite(options.min_width) and math.isfinite(options.max_width)):
        raise ValueError("Widths must be finite")
    if not 0 < options.min_width <= options.max_width <= 255:
        raise ValueError("Column widths must satisfy 0 < min <= max <= 255")


def _text(value: str) -> str:
    if _XML_INVALID.search(value):
        raise ValueError("Text contains characters forbidden by XML 1.0")
    if len(value.encode("utf-16-le")) // 2 > 32767:
        raise WorkbookLimitError("Cell text exceeds the XLSX text limit")
    return value


def _value(value: CellValue) -> CellValue:
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, str):
        return _text(value)
    if isinstance(value, int):
        if abs(value) > 999_999_999_999_999:
            raise ValueError("Integers exceeding 15 digits must be supplied as text")
        return value
    if isinstance(value, float):
        if math.isnan(value):
            return None
        if not math.isfinite(value):
            raise ValueError("Infinite numbers cannot be exported")
        return value
    if isinstance(value, datetime) and value.utcoffset() is not None:
        raise ValueError("Timezone-aware datetimes require explicit conversion by the caller")
    if isinstance(value, date):
        return value
    raise TypeError(f"Unsupported cell value type: {type(value).__name__}")


def _table_columns(table: ReportTable, options: ExcelOptions) -> tuple[str, ...]:
    if not isinstance(table.name, str) or not table.name or len(table.name) > 31:
        raise ValueError("Sheet name must contain 1 to 31 characters")
    if _SHEET_INVALID.search(table.name) or table.name.startswith("'") or table.name.endswith("'"):
        raise ValueError("Invalid worksheet name")
    _text(table.name)
    if isinstance(table.columns, (str, bytes)):
        raise TypeError("Columns must be a sequence of names")
    columns = tuple(table.columns)
    if not columns or len(columns) > min(16384, options.limits.max_columns):
        raise WorkbookLimitError("Invalid or excessive column count")
    if any(not isinstance(column, str) or not column for column in columns):
        raise ValueError("Column names must be nonempty strings")
    if len(set(columns)) != len(columns):
        raise ValueError("Duplicate column names")
    if isinstance(table.start_row, bool) or not isinstance(table.start_row, int):
        raise TypeError("start_row must be an integer")
    if not 1 <= table.start_row <= 1048576:
        raise WorkbookLimitError("start_row exceeds XLSX bounds")
    if table.profile not in PROFILE_IDS:
        raise ValueError("Unknown format profile")
    if set(table.formats) - set(columns):
        raise ValueError("Format override refers to an unknown column")
    for fmt in table.formats.values():
        if not isinstance(fmt, str) or not fmt or len(fmt) > 255:
            raise ValueError("Format overrides must contain 1 to 255 characters")
        _text(fmt)
    return columns


def _row_values(row: ReportRow, columns: tuple[str, ...]) -> list[CellValue]:
    if isinstance(row, Mapping):
        if set(row) != set(columns):
            raise ValueError("Mapping row keys must match columns exactly")
        values = [row[column] for column in columns]
    elif isinstance(row, Sequence) and not isinstance(row, (str, bytes)):
        values = list(row)
        if len(values) != len(columns):
            raise ValueError("Row width must match columns exactly")
    else:
        raise TypeError("Rows must be mappings or sequences of supported scalar values")
    return [_value(value) for value in values]


def _assign(cell: Cell, value: CellValue, budget: _Budget) -> None:
    budget.consume(value)
    cell.value = value
    if isinstance(value, str):
        # Explicit XML string type preserves the original text without active formulas.
        cell.data_type = "s"


def _header(
    ws: Worksheet, columns: tuple[str, ...], table: ReportTable, budget: _Budget
) -> list[float]:
    widths = []
    for index, column in enumerate(columns, 1):
        cell = cast(Cell, ws.cell(table.start_row, index))
        _assign(cell, _text(column), budget)
        cell.fill, cell.font, cell.alignment = _HEADER_FILL, _HEADER_FONT, _HEADER_ALIGNMENT
        if budget.options.border:
            cell.border = _BORDER
        widths.append(max(budget.options.min_width, int(len(column) * 1.15) + 2))
    return widths


def _body(ws: Worksheet, table: ReportTable, columns: tuple[str, ...], budget: _Budget) -> None:
    widths = _header(ws, columns, table, budget)
    last_row = table.start_row
    for offset, row in enumerate(table.rows, 1):
        last_row = table.start_row + offset
        if offset > budget.options.limits.max_rows_per_sheet or last_row > 1048576:
            raise WorkbookLimitError("Sheet row budget exceeded")
        for index, value in enumerate(_row_values(row, columns), 1):
            cell = cast(Cell, ws.cell(last_row, index))
            _assign(cell, value, budget)
            fmt = table.formats.get(columns[index - 1])
            if fmt is None:
                fmt = get_profile_format(table.profile, columns[index - 1], value)
            if fmt != "General":
                cell.number_format = fmt
            if budget.options.zebra and offset % 2 == 0:
                cell.fill = _ZEBRA
            if budget.options.border:
                cell.border = _BORDER
            if value is not None:
                widths[index - 1] = max(widths[index - 1], len(str(value)) + 2)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = min(width, budget.options.max_width)
    if budget.options.freeze_header and table.start_row < 1048576:
        ws.freeze_panes = f"A{table.start_row + 1}"
    if budget.options.auto_filter:
        ws.auto_filter.ref = f"A{table.start_row}:{get_column_letter(len(columns))}{last_row}"


def render(tables: Iterable[ReportTable], options: ExcelOptions) -> bytes:
    """Render validated caller data; preserve ordinary Flowlib cell styling rules."""
    _validate_options(options)
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    names: set[str] = set()
    budget = _Budget(options)
    try:
        for table in tables:
            if not isinstance(table, ReportTable):
                raise TypeError("Expected ReportTable")
            columns = _table_columns(table, options)
            if table.name.casefold() in names:
                raise ValueError("Duplicate worksheet name")
            names.add(table.name.casefold())
            if len(names) > options.limits.max_sheets:
                raise WorkbookLimitError("Worksheet budget exceeded")
            _body(workbook.create_sheet(table.name), table, columns, budget)
        if not names:
            raise ValueError("At least one worksheet is required")
        output = BytesIO()
        workbook.save(output)
        if output.tell() > options.limits.max_output_bytes:
            raise WorkbookLimitError("Workbook output budget exceeded")
        payload = normalize_font_sequence(output.getvalue())
        if len(payload) > options.limits.max_output_bytes:
            raise WorkbookLimitError("Workbook output budget exceeded")
        return payload
    finally:
        workbook.close()
