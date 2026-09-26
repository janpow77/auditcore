"""Führt die Originale aus audit_designer unverändert aus und hält Ein-/Ausgaben fest.

Geladen werden die Blob-geprüften Dateien aus ``janpow77/audit_designer`` (Commit
``eff41a4c``) in ein temporäres Paket ``app`` – ``bpmn_analyzer.py``,
``bpmn_export.py`` und ``core/safe_xml.py`` unverändert, ``validate_bpmn_bva``
per AST aus ``api/bpmn.py`` (das Modul selbst braucht FastAPI/SQLAlchemy).
Die Personalkostenberechnung wird mit einem Stellvertreter für das im
audit_designer fehlende ``PersonnelCostRate``-Modell und einer Scheinsitzung
charakterisiert. Alle Eingaben sind synthetisch::

    python tools/capture_legacy.py /pfad/zu/audit_designer tests/fixtures/legacy_observed.json

Benötigt (wie audit_designer ``backend/requirements.txt``): pandas, openpyxl,
reportlab, defusedxml, lxml.
"""

from __future__ import annotations

import argparse
import ast
import base64
import hashlib
import io
import json
import platform
import random
import subprocess
import sys
import tempfile
import types
from pathlib import Path
from typing import Any

COMMIT = "eff41a4ccedab12b9a73bafff41468712459cb2b"
BASE = "backend/app/modules/flowstat"
BLOBS = {
    f"{BASE}/services/bpmn_analyzer.py": "0cfaf87afca87a379ee8a68edbca6fceea3ab944",
    f"{BASE}/services/bpmn_export.py": "ebc5b0ab2f6df964a44727e5b585c827c6ab891c",
    f"{BASE}/api/bpmn.py": "e7853d7b0356d54fd796071062ea640d506e14a8",
    "backend/app/core/safe_xml.py": "1e3ef65342a6fc3c5392e3519ce1fc6d441c634b",
}
BPMN = "http://www.omg.org/spec/BPMN/20100524/MODEL"
FA = "https://flowaudit.de/bpmn/schema/1.0"


def git_blob(raw: bytes) -> str:
    return hashlib.sha1(  # noqa: S324 - Git-Blob-Kennung
        b"blob " + str(len(raw)).encode() + b"\0" + raw, usedforsecurity=False
    ).hexdigest()


def read_source(repo: Path, path: str) -> bytes:
    raw = subprocess.run(["git", "-C", str(repo), "show", f"{COMMIT}:{path}"], capture_output=True, check=True).stdout
    expected = BLOBS[path]
    if expected and git_blob(raw) != expected:
        raise SystemExit(f"{path}: Blob weicht von der festgehaltenen Fassung ab")
    return raw


class _Column:
    def __init__(self, name: str) -> None:
        self.name = name

    def __eq__(self, other: object) -> Any:  # type: ignore[override]
        return (self.name, other)

    def __hash__(self) -> int:
        return hash(self.name)

    def desc(self) -> str:
        return f"{self.name} desc"


class PersonnelCostRate:
    grade = _Column("grade")
    is_active = _Column("is_active")
    year = _Column("year")


RATES = {"A9": 45.5, "A13": 71.25, "E11": 58.0, "E13": 0.0}


class _Query:
    def __init__(self) -> None:
        self.grade: str | None = None

    def filter(self, *conditions: Any) -> _Query:
        for name, value in conditions:
            if name == "grade":
                self.grade = value
        return self

    def order_by(self, *_args: Any) -> _Query:
        return self

    def first(self) -> Any:
        if self.grade in RATES:
            return types.SimpleNamespace(hourly_cost=RATES[self.grade])
        return None


class FakeSession:
    def query(self, _model: Any) -> _Query:
        return _Query()


def install(repo: Path, target: Path) -> dict[str, Any]:
    app = target / "app"
    for package in ("", "core", "modules", "modules/flowstat", "modules/flowstat/services", "modules/flowstat/models"):
        (app / package).mkdir(parents=True, exist_ok=True)
        (app / package / "__init__.py").write_text("")
    for path in ("services/bpmn_analyzer.py", "services/bpmn_export.py"):
        (app / "modules/flowstat" / path).write_bytes(read_source(repo, f"{BASE}/{path}"))
    (app / "core/safe_xml.py").write_bytes(read_source(repo, "backend/app/core/safe_xml.py"))
    sys.path.insert(0, str(target))
    api = read_source(repo, f"{BASE}/api/bpmn.py").decode("utf-8")
    tree = ast.parse(api)
    function = next(
        node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name == "validate_bpmn_bva"
    )
    import xml.etree.ElementTree as ET

    from defusedxml import ElementTree as DefusedET
    from defusedxml.common import DefusedXmlException

    class BpmnValidationResult:
        def __init__(self, valid: bool, errors: list[str], warnings: list[str]) -> None:
            self.valid, self.errors, self.warnings = valid, errors, warnings

    namespace: dict[str, Any] = {
        "ET": ET,
        "DefusedET": DefusedET,
        "DefusedXmlException": DefusedXmlException,
        "BpmnValidationResult": BpmnValidationResult,
    }
    module = ast.Module(body=[*ast.parse("from __future__ import annotations").body, function], type_ignores=[])
    exec(compile(module, f"{BASE}/api/bpmn.py", "exec"), namespace)  # noqa: S102  # nosec B102
    return namespace


def task(tag: str, tid: str, attrs: dict[str, str], inner: str = "", name: str | None = "Aufgabe") -> str:
    attributes = "".join(f' {k}="{v}"' for k, v in attrs.items())
    name_attr = f' name="{name}"' if name is not None else ""
    return f'<bpmn:{tag} id="{tid}"{name_attr}{attributes}>{inner}</bpmn:{tag}>'


def document(body: str, process_attrs: str = "", extra_ns: str = "") -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        f'<bpmn:definitions xmlns:bpmn="{BPMN}" xmlns:flowaudit="{FA}"{extra_ns} '
        'id="Defs" targetNamespace="urn:test">'
        f'<bpmn:process id="P1"{process_attrs}>{body}</bpmn:process></bpmn:definitions>'
    )


def cases() -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    flow = (
        '<bpmn:startEvent id="S"/><bpmn:endEvent id="E"/>'
        '<bpmn:sequenceFlow id="F1" sourceRef="S" targetRef="T1"/>'
        '<bpmn:sequenceFlow id="F2" sourceRef="T1" targetRef="E"/>'
    )
    result.append({"id": "leer", "xml": document(flow)})
    result.append({"id": "minimal", "xml": document(flow + task("task", "T1", {}))})
    units = [
        "Minuten",
        "Stunden",
        "Std",
        "hours",
        "Tage",
        "days",
        "Wochen",
        "weeks",
        "Monate",
        "months",
        "Sekunden",
        "",
        "MIN",
        "Stunde",
    ]
    for index, unit in enumerate(units):
        attrs = {"durationEstimated": "1.5", "costEstimate": "10", "frequency": "3"}
        if unit:
            attrs["durationUnit"] = unit
        result.append({"id": f"einheit-{index}", "xml": document(flow + task("userTask", "T1", attrs))})
    result.append(
        {
            "id": "aliase",
            "xml": document(
                flow
                + task(
                    "userTask",
                    "T1",
                    {
                        "duration": "45",
                        "cost": "12.50",
                        "resource": "Sachbearbeitung",
                        "frequencyPerYear": "24",
                        "personnel_count": "3",
                    },
                )
                + task("task", "T2", {"durationMinutes": "30", "personnel": "2"})
                + task("serviceTask", "T3", {"fs:costEstimate": "7.25", "fs:frequency": "12"}),
                extra_ns=' xmlns:fs="urn:flowstat"',
            ),
        }
    )
    result.append(
        {
            "id": "ungueltige-zahlen",
            "xml": document(
                task(
                    "task",
                    "T1",
                    {"costEstimate": "abc", "effortPersonDays": "x", "frequency": "1.5", "personnelCount": "zwei"},
                )
                + task("task", "T2", {"costEstimate": "", "frequency": "0", "effortPersonDays": "0"})
            ),
        }
    )
    ext = (
        "<bpmn:documentation>Beschreibung.</bpmn:documentation><bpmn:extensionElements>"
        "<flowaudit:rechtsgrundlage>§ 55 BHO; Art. 74 VO (EU) 2021/1060</flowaudit:rechtsgrundlage>"
        "<flowaudit:interneNotiz>Nur über Schwellenwert.</flowaudit:interneNotiz>"
        "</bpmn:extensionElements>"
    )
    result.append({"id": "erweiterungen", "xml": document(task("userTask", "T1", {}, ext) + task("task", "T2", {}))})
    result.append(
        {
            "id": "erweiterungen-attribut-und-alias",
            "xml": document(
                task(
                    "task",
                    "T1",
                    {},
                    '<bpmn:extensionElements><flowaudit:rechtsgrundlage value="Art. 73 CPR"/>'
                    "<flowaudit:notiz>Alt</flowaudit:notiz></bpmn:extensionElements>",
                )
                + task(
                    "task",
                    "T2",
                    {},
                    "<bpmn:extensionElements><flowaudit:rechtsgrundlage>  </flowaudit:rechtsgrundlage>"
                    "<flowaudit:rechtsgrundlage>Zweite</flowaudit:rechtsgrundlage>"
                    "<flowaudit:notiz>A</flowaudit:notiz><flowaudit:interneNotiz>B</flowaudit:interneNotiz>"
                    "</bpmn:extensionElements>",
                )
            ),
        }
    )
    result.append(
        {
            "id": "strukturiert-1.1",
            "xml": document(
                task(
                    "task",
                    "T1",
                    {},
                    '<bpmn:extensionElements><flowaudit:rechtsgrundlage norm="Verordnung (EU) 2021/1060" '
                    'artikel="74" absatz="2">Artikel 74 Absatz 2 der Verordnung (EU) 2021/1060</flowaudit:rechtsgrundlage>'
                    '<flowaudit:kennzeichen typ="vier_augen"/></bpmn:extensionElements>',
                )
            ),
        }
    )
    result.append(
        {
            "id": "alle-typen-und-verschachtelt",
            "xml": document(
                "".join(
                    task(
                        tag,
                        f"X{i}",
                        {
                            "costEstimate": str(i + 1),
                            "processOwner": f"Stelle {i % 3}",
                            "processDepartment": f"Abteilung {i % 2}",
                            "resourcesSystems": "System A" if i % 2 else "System B",
                        },
                    )
                    for i, tag in enumerate(
                        [
                            "receiveTask",
                            "sendTask",
                            "businessRuleTask",
                            "scriptTask",
                            "manualTask",
                            "serviceTask",
                            "userTask",
                            "task",
                        ]
                    )
                )
                + task(
                    "subProcess",
                    "SP",
                    {"costEstimate": "100"},
                    task("task", "SPT", {"costEstimate": "5"}) + task("callActivity", "CA", {"costEstimate": "9"}),
                )
                + task("callActivity", "CA2", {"costEstimate": "11"})
            ),
        }
    )
    result.append(
        {
            "id": "personal",
            "xml": document(
                task("task", "T1", {"personnelGrade": "A13", "personnelCount": "2", "durationEstimated": "90"})
                + task(
                    "task",
                    "T2",
                    {
                        "personnelGrade": "E11",
                        "personnelCount": "1",
                        "durationEstimated": "2",
                        "durationUnit": "Stunden",
                        "costEstimate": "10",
                        "frequency": "4",
                    },
                )
                + task("task", "T3", {"personnelGrade": "X1", "personnelCount": "1", "durationEstimated": "5"})
                + task("task", "T4", {"personnelGrade": "E13", "personnelCount": "1", "durationEstimated": "5"})
                + task("task", "T5", {"personnelGrade": "A9", "personnelCount": "0", "durationEstimated": "5"})
                + task("task", "T6", {"personnelGrade": "A9", "personnelCount": "1"})
            ),
        }
    )
    result.append(
        {
            "id": "esi-prozess",
            "xml": document(
                task("task", "T1", {}),
                ' esiProfile="ESI-2021" esiCoreRequirements="KA1:K1,K2; KA2 ;KA3:;;KA4: K5 , ,K6"',
            ),
        }
    )
    result.append({"id": "esi-nur-profil", "xml": document(task("task", "T1", {}), ' esiProfile="ESI"')})
    result.append(
        {
            "id": "esi-unterprozess",
            "xml": document(task("subProcess", "SP", {"esiCoreRequirements": "KA7:K1"}, task("task", "T1", {}))),
        }
    )
    result.append({"id": "esi-leer", "xml": document(task("task", "T1", {}), ' esiCoreRequirements="  "')})
    # Validierung
    result.append({"id": "kaputt", "xml": "<broken>"})
    result.append({"id": "kaputt-bpmn", "xml": f'<bpmn:startEvent xmlns:bpmn="{BPMN}"><bpmn:endEvent>'})
    result.append({"id": "doctype-ohne-entitaet", "xml": '<!DOCTYPE x><bpmn:definitions xmlns:bpmn="' + BPMN + '"/>'})
    result.append({"id": "falsche-wurzel", "xml": f'<model:process xmlns:model="{BPMN}" id="P"/>'})
    result.append(
        {
            "id": "praefix-model",
            "xml": (
                f'<model:definitions xmlns:model="{BPMN}" targetNamespace="urn:t"><model:process id="P">'
                '<model:startEvent id="S"/><model:userTask id="T" name="Prüfen" duration="45" cost="12.50" '
                'resource="Sachbearbeitung" frequencyPerYear="24" personnel_count="3"/><model:endEvent id="E"/>'
                "</model:process></model:definitions>"
            ),
        }
    )
    gateway = (
        '<bpmn:startEvent id="S"/><bpmn:exclusiveGateway id="G"/><bpmn:parallelGateway id="G2"/>'
        + "".join(f'<bpmn:sequenceFlow id="GF{i}" sourceRef="G" targetRef="E"/>' for i in range(4))
        + "".join(f'<bpmn:sequenceFlow id="PF{i}" sourceRef="G2" targetRef="E"/>' for i in range(3))
        + '<bpmn:sequenceFlow id="B1" targetRef="E"/><bpmn:sequenceFlow sourceRef="S"/>'
        '<bpmn:sequenceFlow id="B3" sourceRef="Nix" targetRef="Nada"/>'
        '<bpmn:task id="T1"/><bpmn:task id="T1" name="  "/><bpmn:endEvent id="E"/>'
    )
    result.append({"id": "validierung-mix", "xml": document(gateway)})
    result.append({"id": "ohne-start-ende", "xml": document(task("task", "T1", {}))})
    result.append({"id": "ohne-prozess", "xml": f'<bpmn:definitions xmlns:bpmn="{BPMN}"/>'})
    rng = random.Random(20260925)  # noqa: S311 - deterministische Testdaten
    tags = [
        "task",
        "userTask",
        "serviceTask",
        "manualTask",
        "scriptTask",
        "businessRuleTask",
        "sendTask",
        "receiveTask",
        "subProcess",
    ]
    for number in range(60):
        body = []
        for index in range(rng.randint(0, 7)):
            attrs: dict[str, str] = {}
            for key, choices in (
                ("costEstimate", ["", "0", "12.5", "100", "3.333", "-5", "1e3"]),
                ("durationEstimated", ["", "15", "1.25", "abc", "0"]),
                ("durationUnit", ["", "Minuten", "Stunden", "Tage", "Wochen", "Monate", "Jahre"]),
                ("effortPersonDays", ["", "0.5", "2", "x"]),
                ("frequency", ["", "1", "12", "250", "2.5"]),
                ("processOwner", ["", "Referat A", "Referat B"]),
                ("processDepartment", ["", "Abt. 1", "Abt. 2"]),
                ("resourcesSystems", ["", "System X", "System Y"]),
                ("personnelGrade", ["", "A9", "A13", "E11", "Z9"]),
                ("personnelCount", ["", "1", "3"]),
            ):
                value = rng.choice(choices)
                if value:
                    attrs[key] = value
            body.append(
                task(
                    rng.choice(tags),
                    f"R{number}_{index}",
                    attrs,
                    name=rng.choice(["Aufgabe", None, " ", "Prüfen & Zahlen <x>"]) if True else None,
                ).replace("<x>", "&lt;x&gt;")
            )
        result.append({"id": f"zufall-{number:02d}", "xml": document("".join(body))})
    return result


def jsonable(value: Any) -> Any:
    if isinstance(value, float) and value != value:
        return "NaN"
    if isinstance(value, dict):
        return {k: jsonable(v) for k, v in value.items()}
    if isinstance(value, list | tuple):
        return [jsonable(v) for v in value]
    return value


def _rgb(value: Any) -> str | None:
    return value if isinstance(value, str) else None


def workbook_snapshot(data: bytes) -> dict[str, Any]:
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(data))
    sheets = {}
    for sheet in book.worksheets:
        cells = [[jsonable(cell.value) for cell in row] for row in sheet.iter_rows()]
        header = []
        for cell in sheet[1]:
            header.append(
                {
                    "fill": _rgb(cell.fill.fgColor.rgb) if cell.fill and cell.fill.fill_type else None,
                    "font_color": _rgb(cell.font.color.rgb) if cell.font and cell.font.color else None,
                    "bold": bool(cell.font.b),
                    "horizontal": cell.alignment.horizontal,
                    "vertical": cell.alignment.vertical,
                    "border": [
                        cell.border.left.style,
                        cell.border.right.style,
                        cell.border.top.style,
                        cell.border.bottom.style,
                    ],
                }
            )
        widths = {key: dim.width for key, dim in sheet.column_dimensions.items() if dim.width}
        sheets[sheet.title] = {"cells": cells, "header": header, "widths": widths}
    return {"order": book.sheetnames, "sheets": sheets}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("audit_designer", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    import openpyxl
    import pandas as pd
    import reportlab
    from reportlab import rl_config

    rl_config.invariant = 1
    with tempfile.TemporaryDirectory() as temporary:
        namespace = install(args.audit_designer, Path(temporary))
        from app.modules.flowstat.services import bpmn_analyzer, bpmn_export

        personnel_module = types.ModuleType("app.modules.flowstat.models.personnel_cost")
        personnel_module.PersonnelCostRate = PersonnelCostRate  # type: ignore[attr-defined]
        records = []
        for case in cases():
            record: dict[str, Any] = {"id": case["id"], "xml": case["xml"]}
            validation = namespace["validate_bpmn_bva"](case["xml"])
            record["bva"] = {"valid": validation.valid, "errors": validation.errors, "warnings": validation.warnings}
            try:
                record["analysis"] = jsonable(bpmn_analyzer.analyze_bpmn(case["xml"], "Testprozess"))
                analyzer = bpmn_analyzer.BpmnAnalyzer(case["xml"])
                record["esi"] = jsonable(analyzer.extract_esi_requirements())
                sys.modules[personnel_module.__name__] = personnel_module
                try:
                    record["analysis_personnel"] = jsonable(
                        bpmn_analyzer.analyze_bpmn(case["xml"], "Testprozess", db_session=FakeSession())
                    )
                finally:
                    del sys.modules[personnel_module.__name__]
                excel = bpmn_export.export_bpmn_to_excel(case["xml"], "Testprozess").getvalue()
                record["excel"] = workbook_snapshot(excel)
                pdf = bpmn_export.export_bpmn_to_pdf(case["xml"], "Testprozess").getvalue()
                record["pdf_sha256"] = hashlib.sha256(pdf).hexdigest()
                if case["id"] in ("erweiterungen", "alle-typen-und-verschachtelt", "leer"):
                    record["pdf_base64"] = base64.b64encode(pdf).decode()
            except Exception as error:  # noqa: BLE001 - Fehler sind Teil des beobachteten Verhaltens
                record["error"] = {"type": type(error).__name__, "message": str(error)}
            records.append(record)
    output = {
        "source": {"repository": "janpow77/audit_designer", "commit": COMMIT, "blobs": BLOBS},
        "environment": {
            "python": platform.python_version(),
            "pandas": pd.__version__,
            "openpyxl": openpyxl.__version__,
            "reportlab": reportlab.Version,
            "reportlab_invariant": True,
        },
        "personnel_rates": RATES,
        "cases": records,
    }
    args.output.write_text(json.dumps(output, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(len(records), "Fälle")


if __name__ == "__main__":
    main()
