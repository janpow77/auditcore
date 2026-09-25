"""Hilfen der Charakterisierung, die auch die Tests brauchen (ohne pandas)."""

from __future__ import annotations

import io
from typing import Any


def _rgb(value: Any) -> str | None:
    return value if isinstance(value, str) else None


def _json(value: Any) -> Any:
    if isinstance(value, float) and value != value:
        return "NaN"
    return value


def workbook_snapshot(data: bytes) -> dict[str, Any]:
    """Zellwerte, Kopfzeilenformat und Spaltenbreiten aller Blätter."""
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(data))
    sheets = {}
    for sheet in book.worksheets:
        cells = [[_json(cell.value) for cell in row] for row in sheet.iter_rows()]
        header = [
            {
                "fill": _rgb(cell.fill.fgColor.rgb) if cell.fill and cell.fill.fill_type else None,
                "font_color": _rgb(cell.font.color.rgb) if cell.font and cell.font.color else None,
                "bold": bool(cell.font.b),
                "horizontal": cell.alignment.horizontal,
                "vertical": cell.alignment.vertical,
                "border": [
                    cell.border.left.style,
                    cell.border.right.style,
                    cell.border.top.style,
                    cell.border.bottom.style,
                ],
            }
            for cell in sheet[1]
        ]
        widths = {key: dim.width for key, dim in sheet.column_dimensions.items() if dim.width}
        sheets[sheet.title] = {"cells": cells, "header": header, "widths": widths}
    return {"order": book.sheetnames, "sheets": sheets}
