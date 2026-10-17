#!/usr/bin/env python3
"""Erzeugt Golden-Testdaten durch Neuberechnung der Excel-Vorlage in LibreOffice.

Für jeden Testfall werden die Eingaben in eine Kopie der Vorlage geschrieben, die
Mappe wird mit LibreOffice ohne Oberfläche neu berechnet, und die Werte aller
Formelzellen des Zielblatts (auf Wunsch weiterer Blätter) werden als JSON abgelegt.
Golden-Dateien entstehen nur mit diesem Werkzeug und werden nie von Hand geändert.

Aufruf (Wurzel des Repositorys):

    python3 packages-js/sampling-templates-core/scripts/golden_recalc.py \
        packages-js/sampling-templates-core/test/cases/03_MUS_1Z.json \
        packages-js/sampling-templates-core/test/golden/03_MUS_1Z.golden.json

Die Vorlage ist nicht Teil des Repositorys. Sie wird gesucht über ``--vorlage``,
sonst im Verzeichnis ``AUDITCORE_SAMPLING_TEMPLATES_DIR``, sonst in
``packages-js/sampling-templates-core/.vorlage/`` (von git ausgeschlossen).

Format der Falldatei::

    [{"name": "...", "blatt": "03_MUS_1Z",
      "eingaben": {"C7": 48500000, "02_Grundgesamtheit!C25": 1000, "C24": null}}]

Zellen ohne Blattangabe beziehen sich auf ``blatt``; ``null`` leert die Zelle
(auch eine Vorgabe der Vorlage). Voraussetzung: ``soffice`` im PATH, ``openpyxl``.
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

import openpyxl

PACKAGE = Path(__file__).resolve().parents[1]
ENVIRONMENT = "AUDITCORE_SAMPLING_TEMPLATES_DIR"
PATTERN = "Stichprobenvorlagen_*.xlsx"
CONVERT = "xlsx:Calc MS Excel 2007 XML"

Case = dict[str, Any]


def find_template(explicit: str | None) -> Path:
    """Return the workbook from ``--vorlage``, the environment or ``.vorlage/``."""
    if explicit:
        return Path(explicit)
    directory = Path(os.environ.get(ENVIRONMENT, PACKAGE / ".vorlage"))
    matches = sorted(directory.glob(PATTERN))
    if not matches:
        raise SystemExit(
            f"Keine Vorlage ({PATTERN}) in {directory}. "
            f"{ENVIRONMENT} setzen oder --vorlage angeben."
        )
    return matches[-1]


def split_reference(reference: str, sheet: str) -> tuple[str, str]:
    """Split ``Blatt!Zelle``; a bare cell belongs to the sheet of the case."""
    if "!" in reference:
        name, cell = reference.split("!", 1)
        return name, cell
    return sheet, reference


def write_cases(template: Path, cases: list[Case], directory: Path) -> list[Path]:
    """Write one workbook per case with the inputs filled in."""
    files = []
    for index, case in enumerate(cases):
        workbook = openpyxl.load_workbook(template)
        for reference, value in case["eingaben"].items():
            sheet, cell = split_reference(reference, case["blatt"])
            workbook[sheet][cell] = value
        target = directory / f"fall{index:03d}.xlsx"
        workbook.save(target)
        files.append(target)
    return files


#: LibreOffice wandelt in einem Aufruf nur eine begrenzte Zahl von Dateien um (gemessen: 246).
BATCH = 200


def recalculate(files: list[Path], directory: Path) -> Path:
    """Recalculate all workbooks headless in batches; own profile against collisions."""
    output = directory / "out"
    profile = (directory / "profile").as_uri()
    for start in range(0, len(files), BATCH):
        batch = files[start : start + BATCH]
        command = [
            "soffice",
            f"-env:UserInstallation={profile}",
            "--headless",
            "--calc",
            "--convert-to",
            CONVERT,
            "--outdir",
            str(output),
            *[str(path) for path in batch],
        ]
        subprocess.run(command, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)  # noqa: S603
        missing = [path.name for path in batch if not (output / path.name).exists()]
        if missing:
            message = f"LibreOffice hat {len(missing)} Datei(en) nicht umgewandelt: {missing[0]} …"
            raise SystemExit(message)
    return output


def formula_cells(template: Path, sheets: list[str]) -> dict[str, list[str]]:
    """Coordinates of all formula cells per sheet in the pristine template."""
    workbook = openpyxl.load_workbook(template)
    found: dict[str, list[str]] = {}
    for sheet in sheets:
        found[sheet] = [
            cell.coordinate
            for row in workbook[sheet].iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
    return found


def read_expected(path: Path, formulas: dict[str, list[str]]) -> dict[str, Any]:
    """Recalculated values of the formula cells; the empty string becomes ``null``."""
    values = openpyxl.load_workbook(path, data_only=True)
    expected: dict[str, Any] = {}
    for sheet, cells in formulas.items():
        for cell in cells:
            value = values[sheet][cell].value
            expected[f"{sheet}!{cell}"] = None if value == "" else value
    return expected


def build(template: Path, cases: list[Case], extra_sheets: list[str]) -> list[Case]:
    """Recalculate all cases and collect the expected values."""
    with tempfile.TemporaryDirectory(prefix="golden_") as name:
        directory = Path(name)
        output = recalculate(write_cases(template, cases, directory), directory)
        result = []
        for index, case in enumerate(cases):
            sheets = list(dict.fromkeys([case["blatt"], *extra_sheets]))
            formulas = formula_cells(template, sheets)
            expected = read_expected(output / f"fall{index:03d}.xlsx", formulas)
            result.append(
                {
                    "name": case.get("name", f"Fall {index}"),
                    "blatt": case["blatt"],
                    "eingaben": case["eingaben"],
                    "erwartet": expected,
                }
            )
        return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Golden-Testdaten aus der Excel-Vorlage erzeugen")
    parser.add_argument("faelle", help="Falldatei (JSON)")
    parser.add_argument("ziel", help="Golden-Datei (JSON)")
    parser.add_argument("--vorlage", help=f"Pfad der Vorlage (sonst {ENVIRONMENT} oder .vorlage/)")
    parser.add_argument(
        "--blatt",
        action="append",
        default=[],
        help="weiteres Blatt, dessen Formelzellen mit ausgelesen werden (mehrfach möglich)",
    )
    arguments = parser.parse_args()
    if shutil.which("soffice") is None:
        raise SystemExit("LibreOffice (soffice) fehlt im PATH.")
    cases = json.loads(Path(arguments.faelle).read_text(encoding="utf-8"))
    result = build(find_template(arguments.vorlage), cases, arguments.blatt)
    text = json.dumps(result, ensure_ascii=False, indent=1, default=str)
    Path(arguments.ziel).write_text(text + "\n", encoding="utf-8")
    print(f"{len(result)} Fälle geschrieben nach {arguments.ziel}")


if __name__ == "__main__":
    main()
