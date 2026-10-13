#!/usr/bin/env python3
"""Misst an LibreOffice, wie die Tabellenkalkulation vergleicht und aufrundet.

Das Werkzeug braucht die Vorlage nicht: Es legt eine Versuchsmappe an
(Vergleiche ``=``, ``<``, ``<=``, ``>``, ``>=``, Differenz und ``ROUNDUP`` um
viele Werte herum), rechnet sie mit LibreOffice neu und schreibt die
Messwerte nach ``test/fixtures/libreoffice-semantik.json``. Der Test
``test/excel.spec.ts`` prüft die Helfer in ``src/excel`` gegen diese Messung.

Aufruf (Wurzel des Repositorys):

    python3 packages-js/sampling-templates-core/scripts/measure_semantics.py
"""

from __future__ import annotations

import json
import math
import subprocess
import tempfile
from pathlib import Path

import openpyxl

PACKAGE = Path(__file__).resolve().parents[1]
TARGET = PACKAGE / "test" / "fixtures" / "libreoffice-semantik.json"
BASES = [1, 3, 31, 100, 10000, 123456, 1e6, 1e9, 8693697600.0, 0.02, 0.3]
OFFSETS = [0.0] + [
    s * m * 10.0**e for e in (-17, -16, -15, -14, -13) for m in (1, 2, 4, 8) for s in (1, -1)
]
ROUND_BASES = [1, 3, 9, 31, 99, 100, 999, 54321, 123456, 1e6, 2e9, 8693697600.0]
ROUND_STEPS = [0.5, 0.9, 0.95, 1.0, 1.05, 1.5]
ROUND_SINGLE = [
    91.2,
    92.0,
    2000000000.5,
    4850000000.3,
    123456789012.25,
    0.5,
    1e-9,
    -1.2,
    -100.000000001,
]


def comparison_rows() -> list[tuple[float, float]]:
    """Pairs ``(a, b)`` with ``b`` a tiny relative step away from ``a``."""
    return [(base, base * (1 + offset)) for base in BASES for offset in OFFSETS]


def rounding_values() -> list[float]:
    """Values around the twelfth significant digit above an integer, plus singles."""
    values = [float(value) for value in ROUND_SINGLE]
    for base in ROUND_BASES:
        unit = 10.0 ** (math.floor(math.log10(base)) - 11)
        values += [base + step * unit for step in ROUND_STEPS]
    return values


def build(path: Path) -> None:
    """Write the experiment workbook."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "V"
    for row, (a, b) in enumerate(comparison_rows(), start=1):
        sheet.cell(row, 1, a)
        sheet.cell(row, 2, b)
        for column, operator in enumerate(("=", "<", "<=", ">", ">="), start=3):
            sheet.cell(row, column, f"=IF(B{row}{operator}A{row},1,0)")
        sheet.cell(row, 8, f"=B{row}-A{row}")
    rounding = workbook.create_sheet("R")
    for row, value in enumerate(rounding_values(), start=1):
        rounding.cell(row, 1, value)
        rounding.cell(row, 2, f"=ROUNDUP(A{row},0)")
        rounding.cell(row, 3, f"=ROUNDUP(A{row},2)")
    workbook.save(path)


def recalculate(path: Path, directory: Path) -> Path:
    """Recalculate with LibreOffice using a private profile."""
    command = [
        "soffice",
        f"-env:UserInstallation={(directory / 'profile').as_uri()}",
        "--headless",
        "--calc",
        "--convert-to",
        "xlsx:Calc MS Excel 2007 XML",
        "--outdir",
        str(directory / "out"),
        str(path),
    ]
    subprocess.run(command, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)  # noqa: S603
    return directory / "out" / path.name


def read(path: Path) -> dict[str, object]:
    """Collect the measured values."""
    workbook = openpyxl.load_workbook(path, data_only=True)
    compare = []
    for row, (a, b) in enumerate(comparison_rows(), start=1):
        cells = [workbook["V"].cell(row, column).value for column in range(3, 9)]
        compare.append(
            {
                "a": a,
                "b": b,
                "eq": cells[0],
                "lt": cells[1],
                "le": cells[2],
                "gt": cells[3],
                "ge": cells[4],
                "minus": cells[5],
            }
        )
    rounding = []
    for row, value in enumerate(rounding_values(), start=1):
        sheet = workbook["R"]
        rounding.append(
            {"x": value, "roundUp0": sheet.cell(row, 2).value, "roundUp2": sheet.cell(row, 3).value}
        )
    version = subprocess.run(["soffice", "--version"], capture_output=True, text=True, check=False)  # noqa: S603, S607
    return {"quelle": version.stdout.strip(), "vergleiche": compare, "aufrunden": rounding}


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="semantik_") as name:
        directory = Path(name)
        source = directory / "versuch.xlsx"
        build(source)
        result = read(recalculate(source, directory))
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    TARGET.write_text(json.dumps(result, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"Messwerte geschrieben nach {TARGET}")

if __name__ == "__main__":
    main()
